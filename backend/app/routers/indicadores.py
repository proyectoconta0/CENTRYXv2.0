import io
from datetime import date
from typing import Optional

from fastapi import APIRouter, Depends
from fastapi.responses import JSONResponse, StreamingResponse
from sqlalchemy.orm import Session

from database import get_db
from app.services import indicadores_service as svc
from app.core.security import get_empresa_id


class UTF8JSONResponse(JSONResponse):
    # Starlette solo agrega "; charset=utf-8" a media types que empiezan con
    # "text/" — para "application/json" lo omite aunque el body sí sea UTF-8,
    # lo que hace que clientes sin soporte del estándar Fetch/JSON (no todos
    # asumen UTF-8 por defecto) decodifiquen mal tildes y emojis.
    media_type = "application/json; charset=utf-8"


router = APIRouter(default_response_class=UTF8JSONResponse)


def _resolver_rango(desde: Optional[date], hasta: Optional[date]):
    if desde and hasta:
        return desde, hasta
    d, h = svc.rango_default()
    return desde or d, hasta or h


@router.get("/resumen")
def resumen(periodo: str = "mes", db: Session = Depends(get_db),
            empresa_id: Optional[int] = Depends(get_empresa_id)):
    if periodo not in ("mes", "año"):
        periodo = "mes"
    return svc.resumen_simple(db, periodo, empresa_id)


@router.get("/rentabilidad")
def rentabilidad(desde: Optional[date] = None, hasta: Optional[date] = None,
                 db: Session = Depends(get_db),
                 empresa_id: Optional[int] = Depends(get_empresa_id)):
    d, h = _resolver_rango(desde, hasta)
    return svc.rentabilidad(db, d, h, empresa_id)


@router.get("/liquidez")
def liquidez(desde: Optional[date] = None, hasta: Optional[date] = None,
             db: Session = Depends(get_db),
             empresa_id: Optional[int] = Depends(get_empresa_id)):
    d, h = _resolver_rango(desde, hasta)
    return svc.liquidez(db, d, h, empresa_id)


@router.get("/cobranza")
def cobranza(desde: Optional[date] = None, hasta: Optional[date] = None,
             db: Session = Depends(get_db),
             empresa_id: Optional[int] = Depends(get_empresa_id)):
    d, h = _resolver_rango(desde, hasta)
    return svc.cobranza(db, d, h, empresa_id)


@router.get("/gastos")
def gastos(desde: Optional[date] = None, hasta: Optional[date] = None,
           db: Session = Depends(get_db),
           empresa_id: Optional[int] = Depends(get_empresa_id)):
    d, h = _resolver_rango(desde, hasta)
    return svc.gastos_kpi(db, d, h, empresa_id)


@router.get("/ventas")
def ventas(desde: Optional[date] = None, hasta: Optional[date] = None,
           db: Session = Depends(get_db),
           empresa_id: Optional[int] = Depends(get_empresa_id)):
    d, h = _resolver_rango(desde, hasta)
    return svc.ventas_kpi(db, d, h, empresa_id)


@router.get("/evolucion")
def evolucion(meses: int = 6, db: Session = Depends(get_db),
              empresa_id: Optional[int] = Depends(get_empresa_id)):
    return {"data": svc.evolucion(db, meses, empresa_id)}


@router.get("/salud-general")
def salud_general(desde: Optional[date] = None, hasta: Optional[date] = None,
                  db: Session = Depends(get_db),
                  empresa_id: Optional[int] = Depends(get_empresa_id)):
    d, h = _resolver_rango(desde, hasta)
    return svc.salud_general(db, d, h, empresa_id)


def _resolver_periodo(periodo: str) -> str:
    return periodo if periodo in ("mes", "año") else "mes"


@router.get("/detalle/ventas")
def detalle_ventas(periodo: str = "mes", db: Session = Depends(get_db),
                   empresa_id: Optional[int] = Depends(get_empresa_id)):
    d, h = svc.rango_simple(_resolver_periodo(periodo))
    return {"data": svc.detalle_ventas(db, d, h, empresa_id), "desde": d, "hasta": h}


@router.get("/detalle/gastos")
def detalle_gastos(periodo: str = "mes", db: Session = Depends(get_db),
                   empresa_id: Optional[int] = Depends(get_empresa_id)):
    d, h = svc.rango_simple(_resolver_periodo(periodo))
    return {"data": svc.detalle_gastos(db, d, h, empresa_id), "desde": d, "hasta": h}


@router.get("/detalle/ganaste")
def detalle_ganaste(periodo: str = "mes", db: Session = Depends(get_db),
                    empresa_id: Optional[int] = Depends(get_empresa_id)):
    d, h = svc.rango_simple(_resolver_periodo(periodo))
    return {**svc.detalle_ganaste(db, d, h, empresa_id), "desde": d, "hasta": h}


@router.get("/detalle/te-deben")
def detalle_te_deben(db: Session = Depends(get_db),
                     empresa_id: Optional[int] = Depends(get_empresa_id)):
    return {"data": svc.detalle_te_deben(db, empresa_id)}


@router.get("/detalle/en-banco")
def detalle_en_banco(db: Session = Depends(get_db),
                     empresa_id: Optional[int] = Depends(get_empresa_id)):
    return {"data": svc.detalle_en_banco(db, empresa_id=empresa_id)}


@router.get("/detalle/debes-pagar")
def detalle_debes_pagar(db: Session = Depends(get_db),
                        empresa_id: Optional[int] = Depends(get_empresa_id)):
    return {"data": svc.detalle_debes_pagar(db, empresa_id)}


@router.get("/exportar")
def exportar(desde: Optional[date] = None, hasta: Optional[date] = None,
             db: Session = Depends(get_db),
             empresa_id: Optional[int] = Depends(get_empresa_id)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    d, h = _resolver_rango(desde, hasta)
    filas = svc.filas_exportar(db, d, h, empresa_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Indicadores KPI"

    hdr_font  = Font(bold=True, color="FFFFFF")
    hdr_fill  = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center")
    alt_fill  = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    ws.cell(row=1, column=1, value=f"Período: {d} a {h}").font = Font(italic=True, color="666666")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)

    HEADERS = ["Indicador", "Valor", "Semáforo", "Valor Período Anterior", "Variación %"]
    WIDTHS  = [32, 18, 14, 22, 14]
    header_row = 3
    for ci, (hd, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=header_row, column=ci, value=hd)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = hdr_align
        ws.column_dimensions[cell.column_letter].width = w

    SEMAFORO_TXT = {"verde": "🟢 Verde", "amarillo": "🟡 Amarillo", "rojo": "🔴 Rojo", "—": "—"}
    for ri, fila in enumerate(filas, header_row + 1):
        valores = [
            fila["indicador"],
            fila["valor"] if fila["valor"] is not None else "Sin datos",
            SEMAFORO_TXT.get(fila["semaforo"], fila["semaforo"]),
            fila["valor_anterior"] if fila["valor_anterior"] is not None else "—",
            fila["variacion_pct"] if fila["variacion_pct"] is not None else "—",
        ]
        for ci, val in enumerate(valores, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            if (ri - header_row) % 2 == 0:
                cell.fill = alt_fill

    ws.freeze_panes = f"A{header_row + 1}"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    hoy = date.today().strftime("%Y%m%d")
    filename = f"Indicadores_KPI_{hoy}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
