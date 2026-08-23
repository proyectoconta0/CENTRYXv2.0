import io
from datetime import date, datetime, time
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from database import get_db
from app.models.auditoria import AuditoriaLog
from app.models.models import Usuario
from app.core.security import require_administrador

router = APIRouter()


def _serialize(log: AuditoriaLog) -> dict:
    return {
        "id":             log.id,
        "usuario_id":     log.usuario_id,
        "usuario_nombre": log.usuario_nombre or "—",
        "modulo":         log.modulo or "—",
        "accion":         log.accion or "—",
        "descripcion":    log.descripcion or "",
        "ip_address":     log.ip_address,
        "fecha_hora":     log.fecha_hora.strftime("%d/%m/%Y %H:%M:%S") if log.fecha_hora else None,
    }


def _filtrar(q, usuario_id: Optional[int], modulo: Optional[str], desde: Optional[date], hasta: Optional[date]):
    if usuario_id:
        q = q.filter(AuditoriaLog.usuario_id == usuario_id)
    if modulo:
        q = q.filter(AuditoriaLog.modulo == modulo)
    if desde:
        q = q.filter(AuditoriaLog.fecha_hora >= datetime.combine(desde, time.min))
    if hasta:
        q = q.filter(AuditoriaLog.fecha_hora <= datetime.combine(hasta, time.max))
    return q


@router.get("")
def listar(
    usuario_id: Optional[int] = None,
    modulo: Optional[str] = None,
    desde: Optional[date] = None,
    hasta: Optional[date] = None,
    page: int = 1,
    per_page: int = 100,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(require_administrador),
):
    q = _filtrar(db.query(AuditoriaLog), usuario_id, modulo, desde, hasta)
    total = q.count()
    rows = (
        q.order_by(AuditoriaLog.fecha_hora.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
        .all()
    )

    usuarios_disponibles = [
        {"id": u.id, "nombre": u.nombre}
        for u in db.query(Usuario).order_by(Usuario.nombre.asc()).all()
    ]
    modulos_disponibles = [
        m[0] for m in db.query(AuditoriaLog.modulo).distinct().order_by(AuditoriaLog.modulo).all() if m[0]
    ]

    return {
        "total":                 total,
        "page":                  page,
        "per_page":              per_page,
        "data":                  [_serialize(r) for r in rows],
        "usuarios_disponibles":  usuarios_disponibles,
        "modulos_disponibles":   modulos_disponibles,
    }


@router.get("/exportar")
def exportar(
    usuario_id: Optional[int] = None,
    modulo: Optional[str] = None,
    desde: Optional[date] = None,
    hasta: Optional[date] = None,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(require_administrador),
):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(500, "Librería openpyxl no instalada. Ejecute: pip install openpyxl")

    q = _filtrar(db.query(AuditoriaLog), usuario_id, modulo, desde, hasta)
    rows = q.order_by(AuditoriaLog.fecha_hora.desc()).limit(5000).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Log de Auditoría"

    header_font  = Font(bold=True, color="FFFFFF")
    header_fill  = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    alt_fill     = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    HEADERS    = ["Fecha/Hora", "Usuario", "Módulo", "Acción", "Descripción", "IP"]
    COL_WIDTHS = [20, 24, 16, 24, 50, 16]

    for col_idx, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 20

    for row_idx, log in enumerate(rows, start=2):
        s = _serialize(log)
        ws.append([s["fecha_hora"], s["usuario_nombre"], s["modulo"], s["accion"], s["descripcion"], s["ip_address"] or "—"])
        if row_idx % 2 == 0:
            for col_idx in range(1, len(HEADERS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = alt_fill

    ws.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Log_Auditoria_{date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
