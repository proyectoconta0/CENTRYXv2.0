import io
from datetime import date, timedelta, datetime
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.models.models import Proveedor, ProveedorGasto, Gasto, Usuario
from app.core.security import get_current_usuario
from app.services.auditoria_service import registrar_log, ip_de
from database import get_db

router = APIRouter()

TIPOS_DOCUMENTO = ["RUC", "DNI", "Carnet de Extranjería"]
ESTADOS = ["Activo", "Inactivo"]

TIPOS_SIN_CXP = "Anticipo de Proveedor"  # mismo criterio de exclusión que Cuentas por Pagar de Gastos


# ── Schemas ──────────────────────────────────────────────────────────────────

class ProveedorCreate(BaseModel):
    tipo_documento: str
    numero_documento: str
    razon_social: str
    direccion: Optional[str] = None
    distrito: Optional[str] = None
    telefono: Optional[str] = None
    email: Optional[str] = None
    contacto_principal: Optional[str] = None
    cargo_contacto: Optional[str] = None
    estado: str = "Activo"
    observaciones: Optional[str] = None
    numero_cuenta: Optional[str] = None


class ProveedorUpdate(BaseModel):
    tipo_documento: Optional[str] = None
    numero_documento: Optional[str] = None
    razon_social: Optional[str] = None
    direccion: Optional[str] = None
    distrito: Optional[str] = None
    telefono: Optional[str] = None
    email: Optional[str] = None
    contacto_principal: Optional[str] = None
    cargo_contacto: Optional[str] = None
    estado: Optional[str] = None
    observaciones: Optional[str] = None
    numero_cuenta: Optional[str] = None


# ── Helpers ──────────────────────────────────────────────────────────────────

def _semaforo(g: Gasto) -> str:
    if g.estado_pago == "Pagado":
        return "pagado"
    if not g.fecha_vencimiento:
        return "sin_fecha"
    dias = (date.today() - g.fecha_vencimiento).days
    if dias <= 0:
        return "verde"
    elif dias <= 15:
        return "amarillo"
    return "rojo"


def _totales_proveedor(db: Session, proveedor_id: int):
    gastos = db.query(Gasto).filter(Gasto.proveedor_id == proveedor_id).all()
    total_comprado = sum(float(g.monto_soles or g.monto or 0) for g in gastos)
    deuda_pendiente = sum(
        float(g.saldo_pendiente or 0)
        for g in gastos
        if g.estado_pago != "Pagado" and (g.tipo_comprobante or "") != TIPOS_SIN_CXP
    )
    return round(total_comprado, 2), round(deuda_pendiente, 2)


def _serialize(p: Proveedor, db: Session) -> dict:
    total_comprado, deuda_pendiente = _totales_proveedor(db, p.id)
    return {
        "id":                  p.id,
        "tipo_documento":      p.tipo_documento or "",
        "numero_documento":    p.numero_documento,
        "razon_social":        p.razon_social,
        "direccion":           p.direccion or "",
        "distrito":            p.distrito or "",
        "telefono":            p.telefono or "",
        "email":               p.email or "",
        "contacto_principal":  p.contacto_principal or "",
        "cargo_contacto":      p.cargo_contacto or "",
        "estado":              p.estado or "Activo",
        "observaciones":       p.observaciones or "",
        "numero_cuenta":       p.numero_cuenta or "",
        "total_comprado":      total_comprado,
        "deuda_pendiente":     deuda_pendiente,
        "created_at":          str(p.created_at) if p.created_at else None,
        "updated_at":          str(p.updated_at) if p.updated_at else None,
        "creado_por":          p.creado_por,
        "creado_en":           p.creado_en.strftime("%d/%m/%Y %H:%M") if p.creado_en else None,
        "modificado_por":      p.modificado_por,
        "modificado_en":       p.modificado_en.strftime("%d/%m/%Y %H:%M") if p.modificado_en else None,
        "metodo_creacion":     p.metodo_creacion or "Manual",
    }


def _primero_ultimo_mes(hoy: date):
    primero = hoy.replace(day=1)
    if hoy.month == 12:
        ultimo = hoy.replace(day=31)
    else:
        ultimo = hoy.replace(month=hoy.month + 1, day=1) - timedelta(days=1)
    return primero, ultimo


# ── Listar ───────────────────────────────────────────────────────────────────

@router.get("")
def listar(
    search:         str = "",
    estado:         str = "",
    tipo_documento: str = "",
    page:           int = 1,
    per_page:       int = 20,
    db: Session = Depends(get_db),
):
    q = db.query(Proveedor)
    if search:
        like = f"%{search}%"
        q = q.filter(
            Proveedor.razon_social.ilike(like) | Proveedor.numero_documento.ilike(like)
        )
    if estado:
        q = q.filter(Proveedor.estado == estado)
    if tipo_documento:
        q = q.filter(Proveedor.tipo_documento == tipo_documento)

    total = q.count()
    rows  = q.order_by(Proveedor.razon_social.asc()).offset((page - 1) * per_page).limit(per_page).all()
    return {
        "total": total, "page": page, "per_page": per_page,
        "data": [_serialize(p, db) for p in rows],
    }


# ── Crear ────────────────────────────────────────────────────────────────────

@router.post("")
def crear(data: ProveedorCreate, http_request: Request, db: Session = Depends(get_db),
          usuario: Usuario = Depends(get_current_usuario)):
    if db.query(Proveedor).filter(Proveedor.numero_documento == data.numero_documento).first():
        raise HTTPException(400, f"Ya existe un proveedor con el documento {data.numero_documento}")
    if data.estado not in ESTADOS:
        raise HTTPException(400, f"Estado inválido. Use uno de: {', '.join(ESTADOS)}")

    hoy = date.today()
    p = Proveedor(
        tipo_documento     = data.tipo_documento,
        numero_documento   = data.numero_documento,
        razon_social       = data.razon_social,
        direccion          = data.direccion,
        distrito           = data.distrito,
        telefono           = data.telefono,
        email              = data.email,
        contacto_principal = data.contacto_principal,
        cargo_contacto     = data.cargo_contacto,
        estado             = data.estado,
        observaciones      = data.observaciones,
        numero_cuenta      = data.numero_cuenta,
        created_at         = hoy,
        updated_at         = hoy,
        creado_por         = usuario.nombre,
        creado_en          = datetime.utcnow(),
        metodo_creacion    = "Manual",
    )
    db.add(p)
    db.commit()
    db.refresh(p)

    registrar_log(
        db, usuario.id, usuario.nombre, "proveedores", "Creó proveedor",
        f"Creó proveedor {p.razon_social}", ip_de(http_request),
    )

    return _serialize(p, db)


# ── KPIs ─────────────────────────────────────────────────────────────────────

@router.get("/resumen-kpis")
def resumen_kpis(db: Session = Depends(get_db)):
    hoy = date.today()
    primero_mes, ultimo_mes = _primero_ultimo_mes(hoy)

    total_activos = db.query(Proveedor).filter(Proveedor.estado == "Activo").count()

    gastos_mes = db.query(Gasto).filter(
        Gasto.proveedor_id.isnot(None),
        Gasto.fecha >= primero_mes,
        Gasto.fecha <= ultimo_mes,
    ).all()
    total_comprado_mes = round(sum(float(g.monto_soles or g.monto or 0) for g in gastos_mes), 2)

    gastos_pend = db.query(Gasto).filter(
        Gasto.proveedor_id.isnot(None),
        Gasto.estado_pago != "Pagado",
        Gasto.tipo_comprobante != TIPOS_SIN_CXP,
    ).all()
    total_por_pagar = round(sum(float(g.saldo_pendiente or 0) for g in gastos_pend), 2)
    proveedores_deuda_vencida = len({
        g.proveedor_id for g in gastos_pend if _semaforo(g) in ("amarillo", "rojo")
    })

    return {
        "total_proveedores_activos":     total_activos,
        "total_comprado_mes":            total_comprado_mes,
        "total_por_pagar":               total_por_pagar,
        "proveedores_con_deuda_vencida": proveedores_deuda_vencida,
    }


# ── Importar desde proveedores_gastos ──────────────────────────────────────

@router.post("/importar-desde-gastos")
def importar_desde_gastos(db: Session = Depends(get_db)):
    hoy = date.today()
    importados = 0
    for pg in db.query(ProveedorGasto).all():
        if not pg.numero_documento or not pg.nombre_proveedor:
            continue
        if db.query(Proveedor).filter(Proveedor.numero_documento == pg.numero_documento).first():
            continue
        db.add(Proveedor(
            tipo_documento=pg.tipo_documento or "RUC",
            numero_documento=pg.numero_documento,
            razon_social=pg.nombre_proveedor,
            estado="Activo",
            created_at=hoy,
            updated_at=hoy,
        ))
        importados += 1
    db.commit()

    vinculados = 0
    for g in db.query(Gasto).filter(Gasto.proveedor_id.is_(None), Gasto.numero_documento.isnot(None)).all():
        prov = db.query(Proveedor).filter(Proveedor.numero_documento == g.numero_documento).first()
        if prov:
            g.proveedor_id = prov.id
            vinculados += 1
    db.commit()

    return {"importados": importados, "gastos_vinculados": vinculados}


# ── Exportar Excel ───────────────────────────────────────────────────────────

@router.get("/exportar")
def exportar(
    estado:         str = "",
    tipo_documento: str = "",
    db: Session = Depends(get_db),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    q = db.query(Proveedor)
    if estado:         q = q.filter(Proveedor.estado == estado)
    if tipo_documento: q = q.filter(Proveedor.tipo_documento == tipo_documento)
    rows = q.order_by(Proveedor.razon_social.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Proveedores"

    hdr_font  = Font(bold=True, color="FFFFFF")
    hdr_fill  = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center")
    alt_fill  = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    HEADERS = [
        "RUC/Doc", "Tipo Doc", "Proveedor", "Dirección", "Teléfono",
        "Email", "Contacto", "Total Comprado", "Deuda Pendiente", "Estado",
    ]
    WIDTHS = [16, 14, 32, 30, 14, 26, 24, 16, 16, 12]

    for ci, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = hdr_align
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 20

    for ri, p in enumerate(rows, 2):
        total_comprado, deuda_pendiente = _totales_proveedor(db, p.id)
        ws.append([
            p.numero_documento,
            p.tipo_documento or "—",
            p.razon_social,
            p.direccion or "—",
            p.telefono or "—",
            p.email or "—",
            p.contacto_principal or "—",
            total_comprado,
            deuda_pendiente,
            p.estado or "Activo",
        ])
        if ri % 2 == 0:
            for ci in range(1, len(HEADERS) + 1):
                ws.cell(row=ri, column=ci).fill = alt_fill

    ws.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    hoy = date.today().strftime("%Y%m%d")
    filename = f"Proveedores_{hoy}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ── Detalle ──────────────────────────────────────────────────────────────────

@router.get("/{proveedor_id}")
def obtener(proveedor_id: int, db: Session = Depends(get_db)):
    p = db.query(Proveedor).filter(Proveedor.id == proveedor_id).first()
    if not p:
        raise HTTPException(404, "Proveedor no encontrado")
    return _serialize(p, db)


@router.get("/{proveedor_id}/historial-compras")
def historial_compras(
    proveedor_id: int,
    fecha_desde:  Optional[date] = None,
    fecha_hasta:  Optional[date] = None,
    db: Session = Depends(get_db),
):
    if not db.query(Proveedor).filter(Proveedor.id == proveedor_id).first():
        raise HTTPException(404, "Proveedor no encontrado")

    q = db.query(Gasto).filter(Gasto.proveedor_id == proveedor_id)
    if fecha_desde: q = q.filter(Gasto.fecha >= fecha_desde)
    if fecha_hasta: q = q.filter(Gasto.fecha <= fecha_hasta)
    rows = q.order_by(Gasto.fecha.desc()).all()

    data = [{
        "id":                 g.id,
        "fecha":              str(g.fecha),
        "categoria":          g.categoria,
        "descripcion":        g.descripcion or "",
        "tipo_comprobante":   g.tipo_comprobante or "",
        "numero_comprobante": g.numero_comprobante or "",
        "monto":              round(float(g.monto_soles or g.monto or 0), 2),
        "estado_pago":        g.estado_pago or "Pendiente",
    } for g in rows]

    return {"total_comprado": round(sum(d["monto"] for d in data), 2), "data": data}


@router.get("/{proveedor_id}/cuentas-por-pagar")
def cuentas_por_pagar(proveedor_id: int, db: Session = Depends(get_db)):
    if not db.query(Proveedor).filter(Proveedor.id == proveedor_id).first():
        raise HTTPException(404, "Proveedor no encontrado")

    rows = db.query(Gasto).filter(
        Gasto.proveedor_id == proveedor_id,
        Gasto.estado_pago != "Pagado",
        Gasto.tipo_comprobante != TIPOS_SIN_CXP,
    ).order_by(Gasto.fecha_vencimiento.asc()).all()

    data = [{
        "id":                 g.id,
        "fecha":              str(g.fecha),
        "tipo_comprobante":   g.tipo_comprobante or "",
        "numero_comprobante": g.numero_comprobante or "",
        "monto":              round(float(g.monto_soles or g.monto or 0), 2),
        "saldo_pendiente":    round(float(g.saldo_pendiente or 0), 2),
        "fecha_vencimiento":  str(g.fecha_vencimiento) if g.fecha_vencimiento else None,
        "semaforo":           _semaforo(g),
    } for g in rows]

    return {"total_deuda": round(sum(d["saldo_pendiente"] for d in data), 2), "data": data}


# ── Actualizar ───────────────────────────────────────────────────────────────

@router.put("/{proveedor_id}")
def actualizar(proveedor_id: int, data: ProveedorUpdate, http_request: Request, db: Session = Depends(get_db),
                usuario: Usuario = Depends(get_current_usuario)):
    p = db.query(Proveedor).filter(Proveedor.id == proveedor_id).first()
    if not p:
        raise HTTPException(404, "Proveedor no encontrado")

    if data.numero_documento and data.numero_documento != p.numero_documento:
        if db.query(Proveedor).filter(
            Proveedor.numero_documento == data.numero_documento,
            Proveedor.id != proveedor_id,
        ).first():
            raise HTTPException(400, f"Ya existe un proveedor con el documento {data.numero_documento}")
    if data.estado is not None and data.estado not in ESTADOS:
        raise HTTPException(400, f"Estado inválido. Use uno de: {', '.join(ESTADOS)}")

    fields = data.model_dump(exclude_unset=True)
    for k, v in fields.items():
        setattr(p, k, v)
    p.updated_at = date.today()
    p.modificado_por = usuario.nombre
    p.modificado_en  = datetime.utcnow()

    db.commit()
    db.refresh(p)

    registrar_log(
        db, usuario.id, usuario.nombre, "proveedores", "Editó proveedor",
        f"Editó proveedor {p.razon_social}", ip_de(http_request),
    )

    return _serialize(p, db)


@router.delete("/{proveedor_id}")
def eliminar(proveedor_id: int, db: Session = Depends(get_db)):
    p = db.query(Proveedor).filter(Proveedor.id == proveedor_id).first()
    if not p:
        raise HTTPException(404, "Proveedor no encontrado")
    db.query(Gasto).filter(Gasto.proveedor_id == proveedor_id).update({Gasto.proveedor_id: None})
    db.delete(p)
    db.commit()
    return {"mensaje": "Proveedor eliminado"}
