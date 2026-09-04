import io
from datetime import date, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.models.models import OrdenServicio, Cliente, Gasto
from app.models.comercial import VentaComercial
from app.core.security import get_empresa_id
from database import get_db

router = APIRouter()

TIPOS_SERVICIO = [
    "Alquiler de andamios", "Venta de andamios", "Reparación de andamios",
    "Venta de piezas", "Capacitación", "Transporte", "Montaje", "Otros",
]
ESTADOS = ["Pendiente", "En Proceso", "Completada", "Cancelada"]

MONTO_GASTO = func.coalesce(Gasto.monto_soles, Gasto.monto, 0)


# ── Correlativo ──────────────────────────────────────────────────────────────

def _proximo_os(db: Session, empresa_id: Optional[int] = None) -> str:
    q = db.query(OrdenServicio.numero_orden).filter(
        OrdenServicio.numero_orden.isnot(None)
    )
    if empresa_id is not None:
        q = q.filter(OrdenServicio.empresa_id == empresa_id)
    rows = q.all()
    max_num = 0
    for (no,) in rows:
        if no and no.upper().startswith("OS-"):
            try:
                n = int(no[3:])
                if n > max_num:
                    max_num = n
            except ValueError:
                pass
    return f"OS-{max_num + 1:04d}"


# ── Schemas ──────────────────────────────────────────────────────────────────

class OrdenCreate(BaseModel):
    numero_orden: Optional[str] = None
    cliente_id: Optional[int] = None
    tipo_servicio: str
    descripcion: Optional[str] = None
    direccion_obra: Optional[str] = None
    distrito: Optional[str] = None
    fecha_inicio: date
    fecha_fin_estimada: date
    fecha_fin_real: Optional[date] = None
    presupuesto: float
    moneda: str = "PEN"
    tipo_cambio: Optional[float] = None
    avance_porcentaje: int = Field(0, ge=0, le=100)
    estado: str = "Pendiente"
    comprobante_id: Optional[int] = None
    observaciones: Optional[str] = None


class OrdenUpdate(BaseModel):
    cliente_id: Optional[int] = None
    tipo_servicio: Optional[str] = None
    descripcion: Optional[str] = None
    direccion_obra: Optional[str] = None
    distrito: Optional[str] = None
    fecha_inicio: Optional[date] = None
    fecha_fin_estimada: Optional[date] = None
    fecha_fin_real: Optional[date] = None
    presupuesto: Optional[float] = None
    moneda: Optional[str] = None
    tipo_cambio: Optional[float] = None
    avance_porcentaje: Optional[int] = Field(None, ge=0, le=100)
    estado: Optional[str] = None
    comprobante_id: Optional[int] = None
    observaciones: Optional[str] = None


# ── Helpers ──────────────────────────────────────────────────────────────────

def _serialize(o: OrdenServicio, cli: Optional[Cliente] = None) -> dict:
    return {
        "id":                  o.id,
        "numero_orden":        o.numero_orden,
        "cliente_id":          o.cliente_id,
        "cliente_nombre":      cli.razon_social if cli else None,
        "cliente_ruc":         cli.ruc if cli else None,
        "tipo_servicio":       o.tipo_servicio,
        "descripcion":         o.descripcion or "",
        "direccion_obra":      o.direccion_obra or "",
        "distrito":            o.distrito or "",
        "fecha_inicio":        str(o.fecha_inicio) if o.fecha_inicio else None,
        "fecha_fin_estimada":  str(o.fecha_fin_estimada) if o.fecha_fin_estimada else None,
        "fecha_fin_real":      str(o.fecha_fin_real) if o.fecha_fin_real else None,
        "presupuesto":         round(float(o.presupuesto), 2) if o.presupuesto is not None else 0.0,
        "moneda":              o.moneda or "PEN",
        "tipo_cambio":         round(float(o.tipo_cambio), 4) if o.tipo_cambio is not None else None,
        "presupuesto_soles":   round(float(o.presupuesto_soles), 2) if o.presupuesto_soles is not None else round(float(o.presupuesto or 0), 2),
        "avance_porcentaje":   o.avance_porcentaje or 0,
        "estado":              o.estado or "Pendiente",
        "comprobante_id":      o.comprobante_id,
        "observaciones":       o.observaciones or "",
        "created_at":          str(o.created_at) if o.created_at else None,
        "updated_at":          str(o.updated_at) if o.updated_at else None,
    }


def _costo_real(db: Session, orden_id: int, empresa_id: Optional[int] = None) -> float:
    q = db.query(func.coalesce(func.sum(MONTO_GASTO), 0)).filter(Gasto.orden_id == orden_id)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    total = q.scalar()
    return round(float(total or 0), 2)


def _primero_ultimo_mes(hoy: date):
    primero = hoy.replace(day=1)
    if hoy.month == 12:
        ultimo = hoy.replace(day=31)
    else:
        ultimo = hoy.replace(month=hoy.month + 1, day=1) - timedelta(days=1)
    return primero, ultimo


# ── Listar ───────────────────────────────────────────────────────────────────

@router.get("")
def listar(
    search:        str = "",
    tipo_servicio: str = "",
    estado:        str = "",
    fecha_desde:   Optional[date] = None,
    fecha_hasta:   Optional[date] = None,
    page:          int = 1,
    per_page:      int = 20,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(OrdenServicio, Cliente).outerjoin(Cliente, OrdenServicio.cliente_id == Cliente.id)
    if empresa_id is not None:
        q = q.filter(OrdenServicio.empresa_id == empresa_id)
    if search:
        like = f"%{search}%"
        q = q.filter(or_(
            OrdenServicio.numero_orden.ilike(like),
            Cliente.razon_social.ilike(like),
        ))
    if tipo_servicio:
        q = q.filter(OrdenServicio.tipo_servicio == tipo_servicio)
    if estado:
        q = q.filter(OrdenServicio.estado == estado)
    if fecha_desde:
        q = q.filter(OrdenServicio.fecha_inicio >= fecha_desde)
    if fecha_hasta:
        q = q.filter(OrdenServicio.fecha_inicio <= fecha_hasta)

    total = q.count()
    rows  = q.order_by(OrdenServicio.fecha_inicio.desc(), OrdenServicio.id.desc()) \
              .offset((page - 1) * per_page).limit(per_page).all()
    return {
        "total": total, "page": page, "per_page": per_page,
        "data": [_serialize(o, cli) for o, cli in rows],
    }


# ── Crear ────────────────────────────────────────────────────────────────────

@router.post("")
def crear(data: OrdenCreate, db: Session = Depends(get_db),
          empresa_id: Optional[int] = Depends(get_empresa_id)):
    if data.estado not in ESTADOS:
        raise HTTPException(400, f"Estado inválido. Use uno de: {', '.join(ESTADOS)}")

    numero = data.numero_orden or _proximo_os(db, empresa_id)
    q_dup = db.query(OrdenServicio).filter(OrdenServicio.numero_orden == numero)
    if empresa_id is not None:
        q_dup = q_dup.filter(OrdenServicio.empresa_id == empresa_id)
    if q_dup.first():
        raise HTTPException(400, f"Ya existe una orden con el número {numero}")

    if data.comprobante_id:
        q_comp = db.query(VentaComercial).filter(VentaComercial.id == data.comprobante_id)
        if empresa_id is not None:
            q_comp = q_comp.filter(VentaComercial.empresa_id == empresa_id)
        if not q_comp.first():
            raise HTTPException(400, "El comprobante vinculado no existe")

    presupuesto_r = round(data.presupuesto, 2)
    moneda        = data.moneda or "PEN"
    tc            = round(data.tipo_cambio, 4) if data.tipo_cambio else None
    presupuesto_soles = round(presupuesto_r * tc, 2) if moneda == "USD" and tc else presupuesto_r

    orden = OrdenServicio(
        numero_orden       = numero,
        cliente_id         = data.cliente_id,
        tipo_servicio      = data.tipo_servicio,
        descripcion        = data.descripcion,
        direccion_obra     = data.direccion_obra,
        distrito           = data.distrito,
        fecha_inicio       = data.fecha_inicio,
        fecha_fin_estimada = data.fecha_fin_estimada,
        fecha_fin_real     = data.fecha_fin_real,
        presupuesto        = presupuesto_r,
        moneda             = moneda,
        tipo_cambio        = tc,
        presupuesto_soles  = presupuesto_soles,
        avance_porcentaje  = data.avance_porcentaje,
        estado             = data.estado,
        comprobante_id     = data.comprobante_id,
        observaciones      = data.observaciones,
        created_at         = date.today(),
        updated_at         = date.today(),
        empresa_id         = empresa_id,
    )
    db.add(orden)
    db.commit()
    db.refresh(orden)
    cli = db.query(Cliente).filter(Cliente.id == orden.cliente_id).first() if orden.cliente_id else None
    return _serialize(orden, cli)


# ── KPIs ─────────────────────────────────────────────────────────────────────

@router.get("/resumen-kpis")
def resumen_kpis(db: Session = Depends(get_db),
                 empresa_id: Optional[int] = Depends(get_empresa_id)):
    hoy = date.today()
    primero_mes, ultimo_mes = _primero_ultimo_mes(hoy)

    def _agg(q):
        cant  = q.count()
        monto = sum(float(o.presupuesto_soles or o.presupuesto or 0) for o in q.all())
        return cant, round(monto, 2)

    def _q_os(extra_filters=None):
        q = db.query(OrdenServicio)
        if empresa_id is not None:
            q = q.filter(OrdenServicio.empresa_id == empresa_id)
        if extra_filters:
            for f in extra_filters:
                q = q.filter(f)
        return q

    pend_cant, pend_monto = _agg(_q_os([OrdenServicio.estado == "Pendiente"]))
    proc_cant, proc_monto = _agg(_q_os([OrdenServicio.estado == "En Proceso"]))
    comp_cant, comp_monto = _agg(_q_os([
        OrdenServicio.estado == "Completada",
        OrdenServicio.fecha_fin_real.isnot(None),
        OrdenServicio.fecha_fin_real >= primero_mes,
        OrdenServicio.fecha_fin_real <= ultimo_mes,
    ]))
    canc_cant = _q_os([OrdenServicio.estado == "Cancelada"]).count()

    return {
        "pendientes":            {"cantidad": pend_cant, "monto": pend_monto},
        "en_proceso":            {"cantidad": proc_cant, "monto": proc_monto},
        "completadas_mes":       {"cantidad": comp_cant, "monto": comp_monto},
        "canceladas":            {"cantidad": canc_cant},
    }


# ── Calendario ───────────────────────────────────────────────────────────────

@router.get("/calendario")
def calendario(
    mes: int = Query(..., ge=1, le=12),
    anio: int = Query(..., alias="año"),
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    primero = date(anio, mes, 1)
    if mes == 12:
        ultimo = date(anio, 12, 31)
    else:
        ultimo = date(anio, mes + 1, 1) - timedelta(days=1)

    # Órdenes cuyo rango [fecha_inicio, fecha_fin_estimada] se solapa con el mes.
    q = db.query(OrdenServicio, Cliente).outerjoin(
        Cliente, OrdenServicio.cliente_id == Cliente.id
    ).filter(
        OrdenServicio.fecha_inicio <= ultimo,
        or_(OrdenServicio.fecha_fin_estimada >= primero, OrdenServicio.fecha_fin_estimada.is_(None)),
    )
    if empresa_id is not None:
        q = q.filter(OrdenServicio.empresa_id == empresa_id)

    rows = q.all()
    return {"data": [_serialize(o, cli) for o, cli in rows]}


# ── Próximo correlativo ──────────────────────────────────────────────────────

@router.get("/proximo-correlativo")
def proximo_correlativo(db: Session = Depends(get_db),
                        empresa_id: Optional[int] = Depends(get_empresa_id)):
    return {"proximo": _proximo_os(db, empresa_id)}


# ── Exportar Excel ───────────────────────────────────────────────────────────

@router.get("/exportar")
def exportar(
    fecha_desde:   Optional[date] = None,
    fecha_hasta:   Optional[date] = None,
    tipo_servicio: str = "",
    estado:        str = "",
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    q = db.query(OrdenServicio, Cliente).outerjoin(Cliente, OrdenServicio.cliente_id == Cliente.id)
    if empresa_id is not None:
        q = q.filter(OrdenServicio.empresa_id == empresa_id)
    if fecha_desde:   q = q.filter(OrdenServicio.fecha_inicio >= fecha_desde)
    if fecha_hasta:   q = q.filter(OrdenServicio.fecha_inicio <= fecha_hasta)
    if tipo_servicio: q = q.filter(OrdenServicio.tipo_servicio == tipo_servicio)
    if estado:        q = q.filter(OrdenServicio.estado == estado)
    rows = q.order_by(OrdenServicio.fecha_inicio.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Órdenes de Servicio"

    hdr_font  = Font(bold=True, color="FFFFFF")
    hdr_fill  = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center")
    alt_fill  = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    HEADERS = [
        "N° Orden", "Cliente", "RUC", "Tipo Servicio", "Dirección Obra",
        "Fecha Inicio", "Fecha Fin Est.", "Fecha Fin Real",
        "Presupuesto", "Costo Real", "Variación", "Avance %", "Estado",
        "N° Comprobante Vinculado",
    ]
    WIDTHS = [14, 30, 14, 24, 30, 14, 14, 14, 16, 16, 16, 10, 14, 22]

    for ci, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = hdr_align
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 20

    comp_ids = {o.comprobante_id for o, _ in rows if o.comprobante_id}
    comps = {}
    if comp_ids:
        q_comps = db.query(VentaComercial).filter(VentaComercial.id.in_(comp_ids))
        if empresa_id is not None:
            q_comps = q_comps.filter(VentaComercial.empresa_id == empresa_id)
        for vc in q_comps.all():
            comps[vc.id] = vc.numero_factura

    for ri, (o, cli) in enumerate(rows, 2):
        costo_real = _costo_real(db, o.id, empresa_id)
        presupuesto_soles = round(float(o.presupuesto_soles or o.presupuesto or 0), 2)
        variacion = round(presupuesto_soles - costo_real, 2)
        ws.append([
            o.numero_orden or "",
            cli.razon_social if cli else "—",
            cli.ruc if cli else "—",
            o.tipo_servicio or "—",
            o.direccion_obra or "—",
            str(o.fecha_inicio) if o.fecha_inicio else "—",
            str(o.fecha_fin_estimada) if o.fecha_fin_estimada else "—",
            str(o.fecha_fin_real) if o.fecha_fin_real else "—",
            presupuesto_soles,
            costo_real,
            variacion,
            o.avance_porcentaje or 0,
            o.estado or "Pendiente",
            comps.get(o.comprobante_id, "—") if o.comprobante_id else "—",
        ])
        if ri % 2 == 0:
            for ci in range(1, len(HEADERS) + 1):
                ws.cell(row=ri, column=ci).fill = alt_fill

    ws.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    hoy = date.today().strftime("%Y%m%d")
    filename = f"Ordenes_Servicio_{hoy}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ── Detalle ──────────────────────────────────────────────────────────────────

@router.get("/{orden_id}")
def obtener(orden_id: int, db: Session = Depends(get_db),
            empresa_id: Optional[int] = Depends(get_empresa_id)):
    q = db.query(OrdenServicio, Cliente).outerjoin(
        Cliente, OrdenServicio.cliente_id == Cliente.id
    ).filter(OrdenServicio.id == orden_id)
    if empresa_id is not None:
        q = q.filter(OrdenServicio.empresa_id == empresa_id)
    row = q.first()
    if not row:
        raise HTTPException(404, "Orden de servicio no encontrada")
    o, cli = row
    result = _serialize(o, cli)

    costo_real = _costo_real(db, orden_id, empresa_id)
    presupuesto_soles = result["presupuesto_soles"]
    result["costo_real"] = costo_real
    result["variacion"]  = round(presupuesto_soles - costo_real, 2)

    result["comprobante"] = None
    if o.comprobante_id:
        q_vc = db.query(VentaComercial).filter(VentaComercial.id == o.comprobante_id)
        if empresa_id is not None:
            q_vc = q_vc.filter(VentaComercial.empresa_id == empresa_id)
        vc = q_vc.first()
        if vc:
            result["comprobante"] = {
                "id":               vc.id,
                "numero_documento": vc.numero_factura,
                "tipo_documento":   vc.tipo_documento,
                "estado_cobranza":  vc.estado_cobranza,
                "precio_venta":     vc.precio_venta,
            }
    return result


@router.get("/{orden_id}/gastos")
def gastos_de_orden(orden_id: int, db: Session = Depends(get_db),
                    empresa_id: Optional[int] = Depends(get_empresa_id)):
    q_os = db.query(OrdenServicio).filter(OrdenServicio.id == orden_id)
    if empresa_id is not None:
        q_os = q_os.filter(OrdenServicio.empresa_id == empresa_id)
    if not q_os.first():
        raise HTTPException(404, "Orden de servicio no encontrada")

    q_gastos = db.query(Gasto).filter(Gasto.orden_id == orden_id)
    if empresa_id is not None:
        q_gastos = q_gastos.filter(Gasto.empresa_id == empresa_id)
    gastos = q_gastos.order_by(Gasto.fecha.desc()).all()

    data = [{
        "id":                 g.id,
        "fecha":              str(g.fecha),
        "categoria":          g.categoria,
        "descripcion":        g.descripcion or "",
        "monto":              round(float(g.monto_soles or g.monto), 2),
        "proveedor":          g.proveedor or "",
        "tipo_comprobante":   g.tipo_comprobante or "",
        "numero_comprobante": g.numero_comprobante or "",
    } for g in gastos]
    return {"total_gastado": round(sum(d["monto"] for d in data), 2), "data": data}


# ── Actualizar ───────────────────────────────────────────────────────────────

@router.put("/{orden_id}")
def actualizar(orden_id: int, data: OrdenUpdate, db: Session = Depends(get_db),
               empresa_id: Optional[int] = Depends(get_empresa_id)):
    q = db.query(OrdenServicio).filter(OrdenServicio.id == orden_id)
    if empresa_id is not None:
        q = q.filter(OrdenServicio.empresa_id == empresa_id)
    orden = q.first()
    if not orden:
        raise HTTPException(404, "Orden de servicio no encontrada")

    if data.estado is not None and data.estado not in ESTADOS:
        raise HTTPException(400, f"Estado inválido. Use uno de: {', '.join(ESTADOS)}")
    if data.comprobante_id:
        q_vc = db.query(VentaComercial).filter(VentaComercial.id == data.comprobante_id)
        if empresa_id is not None:
            q_vc = q_vc.filter(VentaComercial.empresa_id == empresa_id)
        if not q_vc.first():
            raise HTTPException(400, "El comprobante vinculado no existe")

    fields = data.model_dump(exclude_unset=True)

    moneda_final = fields.get("moneda", orden.moneda or "PEN")
    if "presupuesto" in fields or "moneda" in fields or "tipo_cambio" in fields:
        presupuesto_r = round(fields.get("presupuesto", orden.presupuesto), 2)
        tc = fields.get("tipo_cambio", orden.tipo_cambio)
        tc = round(tc, 4) if tc else None
        fields["presupuesto"] = presupuesto_r
        fields["tipo_cambio"] = tc
        fields["presupuesto_soles"] = round(presupuesto_r * tc, 2) if moneda_final == "USD" and tc else presupuesto_r

    # Si se marca como Completada y no se definió fecha_fin_real, se autocompleta con hoy.
    if fields.get("estado") == "Completada" and not fields.get("fecha_fin_real") and not orden.fecha_fin_real:
        fields["fecha_fin_real"] = date.today()

    for k, v in fields.items():
        setattr(orden, k, v)
    orden.updated_at = date.today()

    db.commit()
    db.refresh(orden)
    cli = db.query(Cliente).filter(Cliente.id == orden.cliente_id).first() if orden.cliente_id else None
    return _serialize(orden, cli)


@router.delete("/{orden_id}")
def eliminar(orden_id: int, db: Session = Depends(get_db),
             empresa_id: Optional[int] = Depends(get_empresa_id)):
    q = db.query(OrdenServicio).filter(OrdenServicio.id == orden_id)
    if empresa_id is not None:
        q = q.filter(OrdenServicio.empresa_id == empresa_id)
    orden = q.first()
    if not orden:
        raise HTTPException(404, "Orden de servicio no encontrada")

    q_gastos = db.query(Gasto).filter(Gasto.orden_id == orden_id)
    if empresa_id is not None:
        q_gastos = q_gastos.filter(Gasto.empresa_id == empresa_id)
    q_gastos.update({Gasto.orden_id: None})

    db.delete(orden)
    db.commit()
    return {"mensaje": "Orden de servicio eliminada"}
