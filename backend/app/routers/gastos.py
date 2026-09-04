import io
import re
import shutil
import uuid
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from sqlalchemy import extract, func, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
import pdfplumber

from app.models.models import Gasto, ProveedorGasto, PagoGasto, Proveedor, Usuario, LoteDetraccion, OrdenPago, LoteDetraccionDetalle, OrdenPagoDetalle, PagoGarantia, Garantia, Prestamo, CuotaPrestamo, MovimientoCaja
from app.models.comercial import CuentaBancaria
from app.models.configuracion import ConfiguracionEmpresa
from app.core.security import get_current_usuario, get_empresa_id
from app.services.comprobante_print import construir_payload_impresion
from app.services.conciliacion_service import limpiar_movimiento_sistema_por_pago
from app.services.auditoria_service import registrar_log, ip_de
from database import get_db

router = APIRouter()

UPLOAD_BASE_GASTOS = Path("uploads/gastos_comprobantes")
ALLOWED_EXT_GASTOS  = {".pdf", ".jpg", ".jpeg", ".png"}

TIPOS_CON_PDF_GASTO = {
    "Factura",
    "Nota de Crédito",
    "Nota de Débito",
    "Recibo de Servicios Públicos",
    "Boleta de Venta",
}

MESES_ES = {1:"Ene",2:"Feb",3:"Mar",4:"Abr",5:"May",6:"Jun",
            7:"Jul",8:"Ago",9:"Sep",10:"Oct",11:"Nov",12:"Dic"}

MESES_LARGOS_ES = {1:"Enero",2:"Febrero",3:"Marzo",4:"Abril",5:"Mayo",6:"Junio",
                    7:"Julio",8:"Agosto",9:"Septiembre",10:"Octubre",11:"Noviembre",12:"Diciembre"}

CATEGORIA_NO_AFECTA_UTILIDAD = "Pagos Tributarios"


def _afecta_utilidad(categoria: Optional[str]) -> bool:
    return categoria != CATEGORIA_NO_AFECTA_UTILIDAD


def _serialize(g: Gasto) -> dict:
    moneda       = g.moneda or "PEN"
    monto_soles  = round(float(g.monto_soles), 2) if g.monto_soles is not None else round(float(g.monto), 2)
    return {
        "id":                  g.id,
        "fecha":               str(g.fecha),
        "categoria":           g.categoria,
        "descripcion":         g.descripcion or "",
        "monto":               round(float(g.monto), 2),
        "area":                g.area or "",
        "tipo_comprobante":    g.tipo_comprobante or "",
        "numero_comprobante":  g.numero_comprobante or "",
        "proveedor":           g.proveedor or "",
        "tipo_documento":      g.tipo_documento or "",
        "numero_documento":    g.numero_documento or "",
        "fecha_vencimiento":   str(g.fecha_vencimiento) if g.fecha_vencimiento else None,
        "saldo_pendiente":     round(float(g.saldo_pendiente), 2) if g.saldo_pendiente is not None else None,
        "estado_pago":         g.estado_pago or "Pendiente",
        "es_recurrente":       bool(g.es_recurrente),
        "recurrente_activo":   bool(g.recurrente_activo),
        "recurrente_padre_id": g.recurrente_padre_id,
        "created_at":          str(g.created_at) if g.created_at else None,
        "base_imponible":      round(float(g.base_imponible), 2) if g.base_imponible is not None else None,
        "igv":                 round(float(g.igv), 2) if g.igv is not None else None,
        "moneda":              moneda,
        "tipo_cambio":         round(float(g.tipo_cambio), 4) if g.tipo_cambio is not None else None,
        "monto_original":      round(float(g.monto_original), 2) if g.monto_original is not None else None,
        "monto_soles":         monto_soles,
        "orden_id":            g.orden_id,
        "numero_orden":        g.orden.numero_orden if g.orden_id and g.orden else None,
        "proveedor_id":        g.proveedor_id,
        "tiene_comprobante":   bool(g.comprobante_path),
        "comprobante_nombre":  g.comprobante_nombre,
        "afecta_utilidad":     g.afecta_utilidad if g.afecta_utilidad is not None else True,
        "periodo_mes":         g.periodo_mes,
        "periodo_anio":        g.periodo_anio,
        "observaciones":       g.observaciones or "",
        "tiene_detraccion":        bool(g.tiene_detraccion),
        "tasa_detraccion":         round(float(g.tasa_detraccion), 2) if g.tasa_detraccion is not None else None,
        "monto_detraccion":        round(float(g.monto_detraccion), 2) if g.monto_detraccion is not None else None,
        "monto_neto_pagar":        round(float(g.monto_neto_pagar), 2) if g.monto_neto_pagar is not None else None,
        "concepto_detraccion":     g.concepto_detraccion,
        "ruc_cuenta_detraccion":   g.ruc_cuenta_detraccion,
        "fecha_limite_detraccion": str(g.fecha_limite_detraccion) if g.fecha_limite_detraccion else None,
        "detraccion_depositada":   bool(g.detraccion_depositada),
        "codigo_detraccion":       g.codigo_detraccion,
        "creado_por":          g.creado_por,
        "creado_en":           g.creado_en.strftime("%d/%m/%Y %H:%M") if g.creado_en else None,
        "modificado_por":      g.modificado_por,
        "modificado_en":       g.modificado_en.strftime("%d/%m/%Y %H:%M") if g.modificado_en else None,
        "metodo_creacion":     g.metodo_creacion or "Manual",
    }


TOLERANCIA_REDONDEO = 5.00

COMPROBANTES_CON_IGV = {"Factura", "Recibo de Servicios Públicos", "Boleta de Venta"}

TIPOS_CON_CODIGO_DETRACCION = {"Factura"}


def _calcular_impuestos(monto: float, tipo_comprobante: Optional[str]):
    if tipo_comprobante in COMPROBANTES_CON_IGV:
        base = round(monto / 1.18, 2)
        igv  = round(monto - base, 2)
        return base, igv
    return None, None


def _semaforo_gasto(g: Gasto) -> str:
    if g.estado_pago == "Pagado":
        return "pagado"
    if not g.fecha_vencimiento:
        return "sin_fecha"
    hoy = date.today()
    dias = (hoy - g.fecha_vencimiento).days
    if dias <= 0:
        return "verde"
    elif dias <= 15:
        return "amarillo"
    else:
        return "rojo"


def _serialize_cpp(g: Gasto) -> dict:
    hoy_d     = date.today()
    dias_venc = 0
    if g.fecha_vencimiento and g.estado_pago != "Pagado":
        dias_venc = max(0, (hoy_d - g.fecha_vencimiento).days)
    d = _serialize(g)
    d["semaforo"]     = _semaforo_gasto(g)
    d["dias_vencido"] = dias_venc
    return d


def _recalcular_pago_gasto(db: Session, g: Gasto):
    pagos       = db.query(PagoGasto).filter(PagoGasto.gasto_id == g.id).all()
    total_pag   = sum(float(p.monto_pagado) for p in pagos)
    monto_total = float(g.monto or 0)
    nuevo_saldo = round(monto_total - total_pag, 2)
    if nuevo_saldo <= 0.01:
        g.saldo_pendiente = 0.0
        g.estado_pago     = "Pagado"
    elif total_pag > 0:
        g.saldo_pendiente = nuevo_saldo
        g.estado_pago     = "Pago Parcial"
    else:
        g.saldo_pendiente = monto_total
        g.estado_pago     = "Pendiente"


def _fecha_limite_detraccion(fecha_emision: date) -> date:
    mes, anio = fecha_emision.month + 1, fecha_emision.year
    if mes > 12:
        mes, anio = 1, anio + 1
    return date(anio, mes, 5)


def _calcular_detraccion_gasto(tipo_comprobante: Optional[str], tiene_detraccion: bool, tasa: Optional[float],
                                total: Optional[float], fecha_emision: Optional[date],
                                fecha_limite_manual: Optional[date]):
    if not tiene_detraccion or tipo_comprobante != "Factura" or not tasa or not total:
        return False, None, None, None, None
    monto_det  = round(total * (tasa / 100))
    monto_neto = round(total - monto_det, 2)
    fecha_lim  = fecha_limite_manual or (_fecha_limite_detraccion(fecha_emision) if fecha_emision else None)
    return True, tasa, monto_det, monto_neto, fecha_lim


def _periodo_label(mes: Optional[int], anio: Optional[int]) -> Optional[str]:
    if not mes or not anio:
        return None
    return f"{MESES_LARGOS_ES.get(mes, mes)} {anio}"


def _semaforo_tributario(g: Gasto) -> str:
    if g.estado_pago == "Pagado":
        return "pagado"
    if not g.fecha_vencimiento:
        return "sin_fecha"
    dias = (g.fecha_vencimiento - date.today()).days
    if dias < 3:
        return "rojo"
    elif dias <= 7:
        return "amarillo"
    else:
        return "verde"


def _serialize_tributario(db: Session, g: Gasto) -> dict:
    saldo = float(g.saldo_pendiente) if g.saldo_pendiente is not None else float(g.monto or 0)
    ultimo_pago = None
    if g.estado_pago == "Pagado":
        ultimo_pago = (
            db.query(PagoGasto)
            .filter(PagoGasto.gasto_id == g.id)
            .order_by(PagoGasto.fecha_pago.desc(), PagoGasto.id.desc())
            .first()
        )
    return {
        "id":               g.id,
        "concepto":         g.descripcion or "",
        "periodo_mes":      g.periodo_mes,
        "periodo_anio":     g.periodo_anio,
        "periodo_label":    _periodo_label(g.periodo_mes, g.periodo_anio),
        "monto":            round(float(g.monto), 2),
        "saldo_pendiente":  round(saldo, 2),
        "fecha_limite":     str(g.fecha_vencimiento) if g.fecha_vencimiento else None,
        "estado":           g.estado_pago or "Pendiente",
        "semaforo":         _semaforo_tributario(g),
        "observaciones":    g.observaciones or "",
        "fecha_pago":       str(ultimo_pago.fecha_pago) if ultimo_pago else None,
        "metodo_pago":      ultimo_pago.metodo_pago if ultimo_pago else None,
        "numero_operacion": ultimo_pago.numero_operacion if ultimo_pago else None,
    }


def _upsert_proveedor(db: Session, tipo_doc: Optional[str], num_doc: Optional[str], nombre: Optional[str],
                      empresa_id: Optional[int] = None) -> Optional[int]:
    if not num_doc or not nombre:
        return None

    existing = db.query(ProveedorGasto).filter(ProveedorGasto.numero_documento == num_doc).first()
    if existing:
        if existing.nombre_proveedor != nombre:
            existing.nombre_proveedor = nombre
            existing.updated_at = date.today()
    else:
        db.add(ProveedorGasto(
            tipo_documento=tipo_doc or None,
            numero_documento=num_doc,
            nombre_proveedor=nombre,
            created_at=date.today(),
            updated_at=date.today(),
        ))

    q_prov = db.query(Proveedor).filter(Proveedor.numero_documento == num_doc)
    if empresa_id is not None:
        q_prov = q_prov.filter(Proveedor.empresa_id == empresa_id)
    prov = q_prov.first()
    if prov:
        if prov.razon_social != nombre:
            prov.razon_social = nombre
            prov.updated_at = date.today()
    else:
        prov = Proveedor(
            tipo_documento=tipo_doc or "RUC",
            numero_documento=num_doc,
            razon_social=nombre,
            estado="Activo",
            empresa_id=empresa_id,
            created_at=date.today(),
            updated_at=date.today(),
        )
        db.add(prov)
        db.flush()
    return prov.id


def _proximo_ri(db: Session, empresa_id: Optional[int] = None) -> str:
    q = db.query(Gasto.numero_comprobante).filter(
        Gasto.tipo_comprobante == "Recibo Interno",
        Gasto.numero_comprobante.isnot(None),
    )
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    rows = q.all()
    max_num = 0
    for (nc,) in rows:
        if nc and nc.upper().startswith("RI-"):
            try:
                n = int(nc[3:])
                if n > max_num:
                    max_num = n
            except ValueError:
                pass
    return f"RI-{max_num + 1:04d}"


def _proximo_gb(db: Session, empresa_id: Optional[int] = None) -> str:
    q = db.query(Gasto.numero_comprobante).filter(
        Gasto.tipo_comprobante == "Gastos Bancarios",
        Gasto.numero_comprobante.isnot(None),
    )
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    rows = q.all()
    max_num = 0
    for (nc,) in rows:
        if nc and nc.upper().startswith("GB-"):
            try:
                n = int(nc[3:])
                if n > max_num:
                    max_num = n
            except ValueError:
                pass
    return f"GB-{max_num + 1:04d}"


def _proximo_ap(db: Session, empresa_id: Optional[int] = None) -> str:
    q = db.query(Gasto.numero_comprobante).filter(
        Gasto.tipo_comprobante == "Anticipo de Proveedor",
        Gasto.numero_comprobante.isnot(None),
    )
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    rows = q.all()
    max_num = 0
    for (nc,) in rows:
        if nc and nc.upper().startswith("AP-"):
            try:
                n = int(nc[3:])
                if n > max_num:
                    max_num = n
            except ValueError:
                pass
    return f"AP-{max_num + 1:04d}"


# ── Schemas ──────────────────────────────────────────────────────────────────

class GastoCreate(BaseModel):
    fecha:              date
    categoria:          str
    descripcion:        str = ""
    monto:              float
    area:               str
    es_recurrente:      bool = False
    tipo_comprobante:   Optional[str] = None
    numero_comprobante: Optional[str] = None
    proveedor:          Optional[str] = None
    tipo_documento:     Optional[str] = None
    numero_documento:   Optional[str] = None
    fecha_vencimiento:  Optional[date] = None
    moneda:             str = "PEN"
    tipo_cambio:        Optional[float] = None
    orden_id:           Optional[int] = None
    periodo_mes:        Optional[int] = None
    periodo_anio:       Optional[int] = None
    observaciones:      Optional[str] = None
    tiene_detraccion:        bool = False
    concepto_detraccion:     Optional[str] = None
    tasa_detraccion:         Optional[float] = None
    ruc_cuenta_detraccion:   Optional[str] = None
    fecha_limite_detraccion: Optional[date] = None
    codigo_detraccion:       Optional[str] = None
    archivo_temp:   Optional[str] = None
    archivo_nombre: Optional[str] = None
    metodo_creacion: Optional[str] = None


class GastoUpdate(BaseModel):
    fecha:              date
    categoria:          str
    descripcion:        str = ""
    monto:              float
    area:               str
    es_recurrente:      bool = False
    tipo_comprobante:   Optional[str] = None
    numero_comprobante: Optional[str] = None
    proveedor:          Optional[str] = None
    tipo_documento:     Optional[str] = None
    numero_documento:   Optional[str] = None
    fecha_vencimiento:  Optional[date] = None
    moneda:             str = "PEN"
    tipo_cambio:        Optional[float] = None
    orden_id:           Optional[int] = None
    periodo_mes:        Optional[int] = None
    periodo_anio:       Optional[int] = None
    observaciones:      Optional[str] = None
    tiene_detraccion:        bool = False
    concepto_detraccion:     Optional[str] = None
    tasa_detraccion:         Optional[float] = None
    ruc_cuenta_detraccion:   Optional[str] = None
    fecha_limite_detraccion: Optional[date] = None
    codigo_detraccion:       Optional[str] = None


class PagoGastoCreate(BaseModel):
    monto_pagado:     float
    fecha_pago:       date
    metodo_pago:      str
    banco:            Optional[str] = None
    numero_cuenta:    Optional[str] = None
    numero_cheque:    Optional[str] = None
    numero_operacion: Optional[str] = None


class PagoGastoUpdate(BaseModel):
    monto_pagado:     float
    fecha_pago:       date
    metodo_pago:      str
    banco:            Optional[str] = None
    numero_cuenta:    Optional[str] = None
    numero_cheque:    Optional[str] = None
    numero_operacion: Optional[str] = None


# ── Listar ──────────────────────────────────────────────────────────────────────

@router.get("")
def listar_gastos(
    search:           str            = "",
    categoria:        str            = "",
    area:             str            = "",
    desde:            Optional[date] = None,
    hasta:            Optional[date] = None,
    solo_recurrentes: bool           = False,
    page:             int            = 1,
    per_page:         int            = 20,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    if solo_recurrentes:
        q = q.filter(Gasto.es_recurrente == True)
    if search:
        q = q.filter(Gasto.descripcion.ilike(f"%{search}%"))
    if categoria:
        q = q.filter(Gasto.categoria == categoria)
    if area:
        q = q.filter(Gasto.area == area)
    if desde:
        q = q.filter(Gasto.fecha >= desde)
    if hasta:
        q = q.filter(Gasto.fecha <= hasta)

    total  = q.count()
    gastos = q.order_by(Gasto.fecha.desc(), Gasto.id.desc()).offset((page - 1) * per_page).limit(per_page).all()
    return {"total": total, "page": page, "per_page": per_page, "data": [_serialize(g) for g in gastos]}


# ── Resumen KPIs ────────────────────────────────────────────────────────────────

@router.get("/resumen")
def resumen_gastos(
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    hoy  = date.today()
    mes  = hoy.month
    anio = hoy.year
    mes_ant  = mes  - 1 if mes  > 1 else 12
    anio_ant = anio     if mes  > 1 else anio - 1

    q_mes = db.query(func.sum(Gasto.monto)).filter(
        extract("month", Gasto.fecha) == mes,
        extract("year",  Gasto.fecha) == anio,
        Gasto.afecta_utilidad.isnot(False),
    )
    if empresa_id is not None:
        q_mes = q_mes.filter(Gasto.empresa_id == empresa_id)
    total_mes = float(q_mes.scalar() or 0)

    q_mes_ant = db.query(func.sum(Gasto.monto)).filter(
        extract("month", Gasto.fecha) == mes_ant,
        extract("year",  Gasto.fecha) == anio_ant,
        Gasto.afecta_utilidad.isnot(False),
    )
    if empresa_id is not None:
        q_mes_ant = q_mes_ant.filter(Gasto.empresa_id == empresa_id)
    total_mes_ant = float(q_mes_ant.scalar() or 0)

    q_top = db.query(
        Gasto.categoria,
        func.sum(Gasto.monto).label("total")
    ).filter(
        extract("month", Gasto.fecha) == mes,
        extract("year",  Gasto.fecha) == anio,
        Gasto.afecta_utilidad.isnot(False),
    )
    if empresa_id is not None:
        q_top = q_top.filter(Gasto.empresa_id == empresa_id)
    top = q_top.group_by(Gasto.categoria).order_by(text("total DESC")).first()

    variacion = 0.0
    if total_mes_ant > 0:
        variacion = round(((total_mes - total_mes_ant) / total_mes_ant) * 100, 1)

    q_anio = db.query(func.sum(Gasto.monto)).filter(
        extract("year", Gasto.fecha) == anio,
        Gasto.afecta_utilidad.isnot(False),
    )
    if empresa_id is not None:
        q_anio = q_anio.filter(Gasto.empresa_id == empresa_id)
    total_anio = float(q_anio.scalar() or 0)

    return {
        "total_mes":      round(total_mes, 2),
        "total_mes_ant":  round(total_mes_ant, 2),
        "variacion_pct":  variacion,
        "total_anio":     round(total_anio, 2),
        "gasto_mas_alto": {"categoria": top[0], "monto": round(float(top[1]), 2)} if top else None,
    }


# ── Por categoría (donut) ────────────────────────────────────────────────────────

@router.get("/por-categoria")
def por_categoria(
    mes:  Optional[int] = None,
    anio: Optional[int] = None,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    hoy = date.today()
    m = mes  or hoy.month
    a = anio or hoy.year

    q = db.query(
        Gasto.categoria,
        func.sum(Gasto.monto).label("total")
    ).filter(
        extract("month", Gasto.fecha) == m,
        extract("year",  Gasto.fecha) == a,
        Gasto.afecta_utilidad.isnot(False),
    )
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    rows = q.group_by(Gasto.categoria).order_by(text("total DESC")).all()

    total = sum(float(r[1]) for r in rows) or 1
    return [
        {"categoria": r[0], "monto": round(float(r[1]), 2), "porcentaje": round((float(r[1]) / total) * 100, 1)}
        for r in rows
    ]


# ── Evolución mensual ────────────────────────────────────────────────────────────

@router.get("/evolucion-mensual")
def evolucion_mensual(
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    hoy = date.today()
    resultado = []
    for i in range(5, -1, -1):
        offset = hoy.month - i - 1
        m = offset % 12 + 1
        a = hoy.year + (offset // 12)
        q = db.query(func.sum(Gasto.monto)).filter(
            extract("month", Gasto.fecha) == m,
            extract("year",  Gasto.fecha) == a,
            Gasto.afecta_utilidad.isnot(False),
        )
        if empresa_id is not None:
            q = q.filter(Gasto.empresa_id == empresa_id)
        total = float(q.scalar() or 0)
        resultado.append({"mes": MESES_ES[m], "anio": a, "total": round(total, 2)})
    return resultado


# ── Reporte por área ─────────────────────────────────────────────────────────────

@router.get("/por-area")
def por_area(
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    hoy = date.today()
    q = db.query(
        Gasto.area,
        func.sum(Gasto.monto).label("total")
    ).filter(
        extract("year", Gasto.fecha) == hoy.year,
        Gasto.afecta_utilidad.isnot(False),
    )
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    rows = q.group_by(Gasto.area).order_by(text("total DESC")).all()

    total = sum(float(r[1]) for r in rows) or 1
    return [
        {"area": r[0] or "Sin área", "monto": round(float(r[1]), 2), "porcentaje": round((float(r[1]) / total) * 100, 1)}
        for r in rows
    ]


# ── Próximo correlativo ─────────────────────────────────────────────────────────

@router.get("/proximo-correlativo")
def proximo_correlativo(
    tipo: str = "RI",
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    if tipo == "GB":
        return {"proximo": _proximo_gb(db, empresa_id)}
    if tipo == "AP":
        return {"proximo": _proximo_ap(db, empresa_id)}
    return {"proximo": _proximo_ri(db, empresa_id)}


# ── Exportar Excel ───────────────────────────────────────────────────────────────

@router.get("/exportar")
def exportar_gastos(
    desde:              Optional[date] = None,
    hasta:              Optional[date] = None,
    categoria:          str            = "",
    area:               str            = "",
    tipos_comprobante:  str            = "",
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise HTTPException(500, "Instale openpyxl: pip install openpyxl")

    q = db.query(Gasto)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    if desde:     q = q.filter(Gasto.fecha >= desde)
    if hasta:     q = q.filter(Gasto.fecha <= hasta)
    if categoria: q = q.filter(Gasto.categoria == categoria)
    if area:      q = q.filter(Gasto.area == area)
    if tipos_comprobante:
        lista_tc = [t.strip() for t in tipos_comprobante.split(",") if t.strip()]
        if lista_tc:
            q = q.filter(Gasto.tipo_comprobante.in_(lista_tc))
    rows = q.order_by(Gasto.fecha.desc()).all()

    wb  = Workbook()
    ws  = wb.active
    ws.title = "Gastos"

    hdr_font  = Font(bold=True, color="FFFFFF")
    hdr_fill  = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center")
    alt_fill  = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    HEADERS = ["Tipo Comprobante", "N° Comprobante", "Proveedor",
               "Moneda", "Monto Original", "T/C", "Base Imponible (S/)", "IGV 18% (S/)", "Monto (S/)", "Área"]
    WIDTHS  = [26, 16, 28, 10, 16, 10, 18, 14, 14, 16]

    for ci, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = hdr_align
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 20

    CON_IGV = {"Factura", "Recibo de Servicios Públicos"}

    total_operativos  = 0.0
    total_tributarios = 0.0

    ri = 2
    for g in rows:
        moneda_g = g.moneda or "PEN"
        monto_s  = round(float(g.monto_soles or g.monto), 2)
        if moneda_g == "USD":
            label_moneda = "US$"
            monto_o      = round(float(g.monto_original), 2) if g.monto_original is not None else monto_s
            tc_g         = round(float(g.tipo_cambio), 4) if g.tipo_cambio is not None else None
        else:
            label_moneda = "S/"
            monto_o      = monto_s
            tc_g         = 1.00
        tiene_igv = g.tipo_comprobante in CON_IGV
        if g.afecta_utilidad is False:
            total_tributarios += monto_s
        else:
            total_operativos += monto_s
        ws.append([
            g.tipo_comprobante or "",
            g.numero_comprobante or "",
            g.proveedor or "",
            label_moneda,
            monto_o,
            tc_g,
            round(float(g.base_imponible), 2) if tiene_igv and g.base_imponible is not None else None,
            round(float(g.igv), 2)             if tiene_igv and g.igv is not None             else None,
            monto_s,
            g.area or "",
        ])
        if ri % 2 == 0:
            for ci in range(1, len(HEADERS) + 1):
                ws.cell(row=ri, column=ci).fill = alt_fill
        ri += 1

    ri += 1
    total_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    total_font = Font(bold=True)
    resumen_filas = [
        ("GASTOS OPERATIVOS",     round(total_operativos, 2)),
        ("PAGOS TRIBUTARIOS",     round(total_tributarios, 2)),
        ("TOTAL SALIDAS DEL MES", round(total_operativos + total_tributarios, 2)),
    ]
    for etiqueta, valor in resumen_filas:
        cell_label = ws.cell(row=ri, column=1, value=etiqueta)
        cell_valor = ws.cell(row=ri, column=9, value=valor)
        cell_label.font = total_font
        cell_valor.font = total_font
        cell_label.fill = total_fill
        cell_valor.fill = total_fill
        ri += 1

    ws.freeze_panes = "A2"
    output = io.BytesIO()
    wb.save(output); output.seek(0)
    fname = f"Reporte_Gastos_{date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


# ── Crear gasto ──────────────────────────────────────────────────────────────────

@router.post("")
def crear_gasto(data: GastoCreate, http_request: Request, db: Session = Depends(get_db),
                 usuario: Usuario = Depends(get_current_usuario),
                 empresa_id: Optional[int] = Depends(get_empresa_id)):
    return _crear_gasto(data, http_request, db, usuario, empresa_id)


def _crear_gasto(data: GastoCreate, http_request: Request, db: Session, usuario: Usuario,
                 empresa_id: Optional[int] = None) -> dict:
    if data.monto <= 0:
        raise HTTPException(400, "El monto debe ser mayor a 0")

    numero = data.numero_comprobante
    if data.tipo_comprobante == "Recibo Interno" and not numero:
        numero = _proximo_ri(db, empresa_id)
    elif data.tipo_comprobante == "Gastos Bancarios" and not numero:
        numero = _proximo_gb(db, empresa_id)
    elif data.tipo_comprobante == "Anticipo de Proveedor" and not numero:
        numero = _proximo_ap(db, empresa_id)

    es_anticipo = data.tipo_comprobante == "Anticipo de Proveedor"

    monto_r   = round(data.monto, 2)
    moneda    = data.moneda or "PEN"
    tc        = round(data.tipo_cambio, 4) if data.tipo_cambio else None
    monto_orig = monto_r if moneda == "USD" else None
    m_soles    = round(monto_r * tc, 2) if moneda == "USD" and tc else monto_r
    base, igv  = _calcular_impuestos(m_soles, data.tipo_comprobante)
    monto_bd   = m_soles

    tiene_det, tasa_det, monto_det, monto_neto, fecha_lim_det = _calcular_detraccion_gasto(
        data.tipo_comprobante, data.tiene_detraccion, data.tasa_detraccion, monto_bd,
        data.fecha, data.fecha_limite_detraccion,
    )

    g = Gasto(
        fecha=data.fecha,
        categoria=data.categoria,
        descripcion=data.descripcion.strip(),
        monto=monto_bd,
        area=data.area,
        es_recurrente=data.es_recurrente,
        recurrente_activo=data.es_recurrente,
        tipo_comprobante=data.tipo_comprobante or None,
        numero_comprobante=numero or None,
        proveedor=data.proveedor or None,
        tipo_documento=data.tipo_documento or None,
        numero_documento=data.numero_documento or None,
        fecha_vencimiento=data.fecha_vencimiento or None,
        saldo_pendiente=None if es_anticipo else monto_bd,
        estado_pago=None if es_anticipo else "Pendiente",
        created_at=date.today(),
        base_imponible=base,
        igv=igv,
        moneda=moneda,
        tipo_cambio=tc,
        monto_original=monto_orig,
        monto_soles=m_soles,
        orden_id=data.orden_id,
        afecta_utilidad=_afecta_utilidad(data.categoria),
        periodo_mes=data.periodo_mes,
        periodo_anio=data.periodo_anio,
        observaciones=data.observaciones or None,
        tiene_detraccion=tiene_det,
        tasa_detraccion=tasa_det,
        monto_detraccion=monto_det,
        monto_neto_pagar=monto_neto,
        concepto_detraccion=data.concepto_detraccion if tiene_det else None,
        ruc_cuenta_detraccion=data.ruc_cuenta_detraccion if tiene_det else None,
        fecha_limite_detraccion=fecha_lim_det,
        codigo_detraccion=data.codigo_detraccion if data.tipo_comprobante in TIPOS_CON_CODIGO_DETRACCION else None,
        creado_por=usuario.nombre,
        creado_en=datetime.utcnow(),
        metodo_creacion=data.metodo_creacion or ("Importación ZIP" if data.archivo_temp else "Manual"),
        empresa_id=empresa_id,
    )
    db.add(g)
    g.proveedor_id = _upsert_proveedor(db, data.tipo_documento, data.numero_documento, data.proveedor, empresa_id)
    db.commit()
    db.refresh(g)

    if data.archivo_temp:
        origen = TMP_ZIP_DIR_GASTOS / Path(data.archivo_temp).name
        if origen.is_file():
            if g.tipo_comprobante in TIPOS_CON_PDF_GASTO:
                upload_dir = UPLOAD_BASE_GASTOS / str(g.id)
                upload_dir.mkdir(parents=True, exist_ok=True)
                nombre_final = data.archivo_nombre or origen.name
                destino = upload_dir / nombre_final
                shutil.move(str(origen), str(destino))
                g.comprobante_path = str(destino)
                g.comprobante_nombre = nombre_final
                db.commit()
            else:
                origen.unlink(missing_ok=True)

    registrar_log(
        db, usuario.id, usuario.nombre, "gastos", "Registró gasto",
        f"Registró gasto S/ {float(g.monto or 0):,.2f} - categoría {g.categoria}",
        ip_de(http_request),
    )

    return _serialize(g)

# ── Importación de PDF (facturas/boletas de proveedor) ─────────────────────────

TAMANO_MAXIMO_PDF_GASTO = 10 * 1024 * 1024
TAMANO_MAXIMO_ZIP_GASTO = 50 * 1024 * 1024
MAX_PDFS_POR_ZIP_GASTO  = 100

TMP_ZIP_DIR_GASTOS = Path("uploads/gastos_comprobantes/_tmp_zip")

_PATRONES_TIPO_GASTO = [
    ("Recibo por Honorarios", r'RECIBO\s+POR\s+HONORARIOS\s+ELECTR[OÓ]NICO'),
    ("Boleta de Venta",       r'BOLETA\s+DE\s+VENTA'),
    ("Factura",               r'FACTURA\s+ELECTR[OÓ]NICA'),
]
_PREFIJOS_NUMERO_GASTO = {
    "Boleta de Venta":       [r'\b(B\d{3}\s*[-–]\s*\d{1,8})\b'],
    "Factura":               [r'\b([FE]\d{3}\s*[-–]\s*\d{1,8})\b'],
    "Recibo por Honorarios": [r'\b(E\d{3}\s*[-–]\s*\d{1,8})\b'],
}
_PATRON_NUMERO_DOC_GASTO = r'\b([A-Z]{1,4}\d{0,4}\s*[-–]\s*\d{1,8})\b'

_CAMPOS_CLAVE_GASTO = [
    "numero_comprobante", "tipo_comprobante", "numero_documento",
    "proveedor", "fecha", "monto",
]


def _ruc_empresa_propia(db: Session, empresa_id: Optional[int] = None) -> Optional[str]:
    q = db.query(ConfiguracionEmpresa)
    if empresa_id is not None:
        q = q.filter(ConfiguracionEmpresa.id == empresa_id)
    cfg = q.first()
    return cfg.ruc if cfg and cfg.ruc else None


def _detectar_numero_y_tipo_gasto(texto: str):
    tipo = None
    for tipo_candidato, patron_tipo in _PATRONES_TIPO_GASTO:
        if re.search(patron_tipo, texto, re.IGNORECASE):
            tipo = tipo_candidato
            break

    numero = None
    for patron in _PREFIJOS_NUMERO_GASTO.get(tipo, []):
        m = re.search(patron, texto, re.IGNORECASE)
        if m:
            numero = re.sub(r'\s+', '', m.group(1)).upper().replace('–', '-')
            break
    if not numero:
        m = re.search(_PATRON_NUMERO_DOC_GASTO, texto, re.IGNORECASE)
        if m:
            numero = re.sub(r'\s+', '', m.group(1)).upper().replace('–', '-')

    return tipo, numero


_SUFIJO_SOCIETARIO = r'(?:S\.?\s?A\.?\s?C\.?|S\.?\s?A\.?\s?A\.?|S\.?\s?R\.?\s?L\.?|E\.?\s?I\.?\s?R\.?\s?L\.?|S\.?\s?A\.?)'


def _extraer_razon_social_con_sufijo(segmento: str) -> Optional[str]:
    matches = list(re.finditer(
        r'([A-ZÁÉÍÓÚÑ0-9][A-ZÁÉÍÓÚÑ0-9\s&,\.\-]{2,100}?' + _SUFIJO_SOCIETARIO + r')(?=\s|,|\n|$)',
        segmento, re.IGNORECASE,
    ))
    if not matches:
        return None
    return re.sub(r'\s{2,}', ' ', matches[-1].group(1)).strip()[:120] or None


def _extraer_emisor_gasto(texto: str, ruc_propio: Optional[str]):
    matches = list(re.finditer(r'RUC\s*:?\s*(\d{11})', texto, re.IGNORECASE))
    if not matches:
        return None, None

    candidatos = [m for m in matches if m.group(1) != ruc_propio] if ruc_propio else matches
    if not candidatos:
        candidatos = matches

    elegido = None
    for m in candidatos:
        ventana_previa = texto[max(0, m.start() - 60): m.start()]
        if not re.search(r'SE[NÑ]OR|CLIENTE|ADQUIRIENTE', ventana_previa, re.IGNORECASE):
            elegido = m
            break
    if elegido is None:
        elegido = candidatos[0]

    ruc = elegido.group(1)

    inicio_linea = texto.rfind("\n", 0, elegido.start())
    contexto_previo = texto[max(0, inicio_linea - 200): elegido.start()]
    razon_social = _extraer_razon_social_con_sufijo(contexto_previo)

    if not razon_social:
        lineas_previas = [l.strip() for l in contexto_previo.split("\n") if l.strip()]
        if lineas_previas:
            razon_social = re.sub(r'\s{2,}', ' ', lineas_previas[-1]).strip()[:120] or None

    return ruc, razon_social


def _extraer_emisor_recibo_honorarios(texto: str, ruc_propio: Optional[str]):
    m = re.search(r'RECIBO\s+POR\s+HONORARIOS\s+ELECTR[OÓ]NICO', texto, re.IGNORECASE)
    if not m:
        return None, None

    lineas_previas = [l.strip() for l in texto[:m.start()].split("\n") if l.strip()]
    for i in range(len(lineas_previas) - 1, -1, -1):
        if re.fullmatch(r'\d{11}', lineas_previas[i]) and lineas_previas[i] != ruc_propio:
            nombre = lineas_previas[i + 1] if i + 1 < len(lineas_previas) else None
            if nombre:
                nombre = re.sub(r'\s{2,}', ' ', nombre).strip()[:120] or None
            return lineas_previas[i], nombre
    return None, None


def _extraer_ruc_recibo_honorarios_gasto(texto: str, ruc_propio: Optional[str]) -> Optional[str]:
    for m in re.finditer(r'R\.?U\.?C\.?\s*:?\s*(1[0-9]{10})', texto, re.IGNORECASE):
        if m.group(1) != ruc_propio:
            return m.group(1)
    return None


def _extraer_ruc_persona_natural_gasto(texto: str, ruc_propio: Optional[str],
                                        limite: Optional[int] = None) -> Optional[str]:
    if limite is not None:
        for m in re.finditer(r'\b(10\d{9})\b', texto[:limite]):
            if m.group(1) != ruc_propio:
                return m.group(1)
    for m in re.finditer(r'\b(10\d{9})\b', texto):
        if m.group(1) != ruc_propio:
            return m.group(1)
    return None


def _extraer_fecha_gasto(texto: str, patron: str):
    m = re.search(patron, texto, re.IGNORECASE)
    if not m:
        return None
    partes = re.split(r'[/-]', m.group(1))
    if len(partes) != 3:
        return None
    d, mes, anio = partes
    try:
        return date(int(anio), int(mes), int(d)).isoformat()
    except ValueError:
        return None


_MESES_ES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}


def _extraer_fecha_texto_gasto(texto: str, patron: str):
    m = re.search(patron, texto, re.IGNORECASE)
    if not m:
        return None
    mes = _MESES_ES.get(m.group(2).lower())
    if not mes:
        return None
    try:
        return date(int(m.group(3)), mes, int(m.group(1))).isoformat()
    except ValueError:
        return None


def _extraer_fecha_libre_es_gasto(texto: str) -> Optional[str]:
    patron = r'(\d{1,2})\s+de\s+(' + '|'.join(_MESES_ES.keys()) + r')\s+del?\s+(\d{4})'
    m = re.search(patron, texto, re.IGNORECASE)
    if not m:
        return None
    mes = _MESES_ES.get(m.group(2).lower())
    if not mes:
        return None
    try:
        return date(int(m.group(3)), mes, int(m.group(1))).isoformat()
    except ValueError:
        return None


def _extraer_monto_gasto(texto: str, patron: str):
    m = re.search(patron, texto, re.IGNORECASE)
    if not m:
        return None
    crudo = m.group(1).replace(",", "")
    try:
        return round(float(crudo), 2)
    except ValueError:
        return None


def _extraer_total_gasto(texto: str) -> Optional[float]:
    m = re.search(r'IMPORTE\s+TOTAL\s*:?\s*S?/?\.?\s*([\d,]+\.\d{2})', texto, re.IGNORECASE)
    if not m:
        for candidata in re.finditer(r'\bTOTAL\s*:?\s*S?/?\.?\s*([\d,]+\.\d{2})', texto, re.IGNORECASE):
            precontexto = texto[max(0, candidata.start() - 5): candidata.start()]
            if re.search(r'SUB[\s-]?$', precontexto, re.IGNORECASE):
                continue
            m = candidata
    if not m:
        return None
    crudo = m.group(1).replace(",", "")
    try:
        return round(float(crudo), 2)
    except ValueError:
        return None


def _extraer_total_honorarios_gasto(texto: str) -> Optional[float]:
    return _extraer_monto_gasto(texto, r'TOTAL\s+POR\s+HONORARIOS\s*:?\s*S?/?\.?\s*([\d,]+\.\d{2})')


def _extraer_forma_pago_gasto(texto: str) -> Optional[str]:
    m = re.search(r'FORMA\s+DE\s+PAGO\s*:?\s*([^\n]+)', texto, re.IGNORECASE)
    if not m:
        return None
    return re.sub(r'\s{2,}', ' ', m.group(1)).strip()[:100] or None


def _extraer_descripcion_gasto(texto: str):
    m = re.search(r'DESCRIPCI[OÓ]N\s*\n\s*([^\n]+)', texto, re.IGNORECASE)
    if m:
        return re.sub(r'\s{2,}', ' ', m.group(1)).strip()[:200] or None

    lineas = texto.split("\n")
    for i, linea in enumerate(lineas):
        if not re.search(r'DESCRIPCI[OÓ]N', linea, re.IGNORECASE):
            continue
        for candidata in lineas[i + 1: i + 4]:
            candidata = candidata.strip()
            if not candidata:
                continue
            fila = re.match(r'^[\d.,]+\s+\S+\s+(.+?)\s+[\d.,]+(?:\s+[\d.,]+)*$', candidata)
            texto_desc = fila.group(1) if fila else candidata
            return re.sub(r'\s{2,}', ' ', texto_desc).strip()[:200] or None
        return None

    m = re.search(r'POR\s+CONCEPTO\s+DE\s*:?\s*\n?\s*([^\n]+)', texto, re.IGNORECASE)
    if m:
        return re.sub(r'\s{2,}', ' ', m.group(1)).strip()[:200] or None

    m = re.search(r'^(.*SERVICIO\s+DE\s+.+)$', texto, re.IGNORECASE | re.MULTILINE)
    if m:
        return re.sub(r'\s{2,}', ' ', m.group(1)).strip()[:200] or None

    return None


def _extraer_datos_pdf_gasto(texto: str, ruc_propio: Optional[str]) -> dict:
    tipo_comprobante, numero_comprobante = _detectar_numero_y_tipo_gasto(texto)
    es_honorarios = tipo_comprobante == "Recibo por Honorarios"

    ruc_proveedor, proveedor = _extraer_emisor_gasto(texto, ruc_propio)
    if es_honorarios and not ruc_proveedor:
        ruc_proveedor, proveedor = _extraer_emisor_recibo_honorarios(texto, ruc_propio)
    if es_honorarios and not ruc_proveedor:
        ruc_proveedor = _extraer_ruc_recibo_honorarios_gasto(texto, ruc_propio)
    if es_honorarios and not ruc_proveedor:
        limite = None
        if numero_comprobante:
            m_num = re.search(re.escape(numero_comprobante), texto, re.IGNORECASE)
            if m_num:
                limite = m_num.start()
        ruc_proveedor = _extraer_ruc_persona_natural_gasto(texto, ruc_propio, limite)

    if es_honorarios:
        fecha = _extraer_fecha_texto_gasto(
            texto, r'FECHA\s+DE\s+EMISI[OÓ]N\s*:?\s*(\d{1,2})\s+(?:DE\s+)?(\w+)\s+(?:DE\s+)?(\d{4})',
        )
        if not fecha:
            fecha = _extraer_fecha_libre_es_gasto(texto)
    else:
        fecha = _extraer_fecha_gasto(texto, r'FECHA\s*(?:DE\s+)?EMISI[OÓ]N\s*:?\s*(\d{2}[/-]\d{2}[/-]\d{4})')
    fecha_vencimiento  = _extraer_fecha_gasto(texto, r'FECHA\s*(?:DE\s+)?VENCIMIENTO\s*:?\s*(\d{2}[/-]\d{2}[/-]\d{4})')

    retencion_monto = None
    total_neto      = None
    if es_honorarios:
        monto          = _extraer_total_honorarios_gasto(texto) or _extraer_total_gasto(texto)
        base_imponible = monto
        igv            = 0.0
        if monto is not None:
            retencion_monto = 0.0
            total_neto      = monto
    else:
        base_imponible = _extraer_monto_gasto(texto, r'(?:OP\.?\s*GRAVADAS?|BASE\s+IMPONIBLE|VALOR\s+VENTA|SUB[\s-]?TOTAL(?:\s+VENTAS)?)[:\s]+S?/?\.?\s*([\d,]+\.\d{2})')
        igv            = _extraer_monto_gasto(texto, r'I[GU]V\s*(?:1[08]\.?0?\s?%)?\s*:?\s*S?/?\.?\s*([\d,]+\.\d{2})')
        monto          = _extraer_total_gasto(texto)
        if base_imponible is None and monto is not None and igv is not None:
            base_imponible = round(monto - igv, 2)

    descripcion = _extraer_descripcion_gasto(texto)

    detectados = {
        "numero_comprobante": numero_comprobante,
        "tipo_comprobante":   tipo_comprobante,
        "numero_documento":   ruc_proveedor,
        "proveedor":          proveedor,
        "fecha":              fecha,
        "monto":              monto,
    }
    campos_no_detectados = [c for c in _CAMPOS_CLAVE_GASTO if not detectados.get(c)]

    if es_honorarios:
        if tipo_comprobante and ruc_proveedor and monto:
            confianza = 1.0
        elif tipo_comprobante and monto:
            confianza = 0.8
        elif tipo_comprobante:
            confianza = 0.5
        else:
            confianza = 0.0
    else:
        confianza = round(1 - len(campos_no_detectados) / len(_CAMPOS_CLAVE_GASTO), 2)

    campos_no_detectados = campos_no_detectados + ["categoria", "area"]

    resultado = {
        "tipo_comprobante":     tipo_comprobante,
        "numero_comprobante":   numero_comprobante,
        "fecha":                fecha,
        "fecha_vencimiento":    fecha_vencimiento,
        "tipo_documento":       "RUC" if ruc_proveedor else None,
        "numero_documento":     ruc_proveedor,
        "proveedor":            proveedor,
        "descripcion":          descripcion,
        "base_imponible":       base_imponible,
        "igv":                  igv,
        "monto":                monto,
        "moneda":               "PEN",
        "confianza":            max(0.0, min(1.0, confianza)),
        "campos_no_detectados": campos_no_detectados,
    }
    if es_honorarios:
        resultado["retencion_monto"]  = retencion_monto
        resultado["total_neto"]       = total_neto
        resultado["forma_pago"]       = _extraer_forma_pago_gasto(texto)
        resultado["tiene_detraccion"] = False
        resultado["detraccion_monto"] = 0
        resultado["codigo_detraccion"] = None
    return resultado


def _procesar_pdf_individual_gasto(contenido: bytes, nombre_archivo: str, db: Session,
                                    ruc_propio: Optional[str],
                                    empresa_id: Optional[int] = None) -> dict:
    try:
        with pdfplumber.open(io.BytesIO(contenido)) as pdf:
            paginas_texto = [pagina.extract_text() or "" for pagina in pdf.pages]
    except Exception:
        return {
            "estado": "no_valido", "no_valido_ilegible": True,
            "error": "No se pudo leer el archivo. Verifique que sea un PDF válido",
        }

    texto = "\n".join(paginas_texto)

    if len(texto.strip()) < 20:
        return {
            "estado": "no_valido", "no_valido_ilegible": True,
            "error": ("Este PDF es una imagen escaneada. Los datos no pudieron extraerse "
                      "automáticamente. Por favor usa Nuevo Gasto para registrarlo."),
        }

    resultado = _extraer_datos_pdf_gasto(texto, ruc_propio)
    resultado["no_valido_ilegible"] = False

    existente = None
    if resultado["numero_comprobante"] and resultado["numero_documento"]:
        q_dup = db.query(Gasto).filter(
            Gasto.numero_comprobante == resultado["numero_comprobante"],
            Gasto.numero_documento == resultado["numero_documento"],
        )
        if empresa_id is not None:
            q_dup = q_dup.filter(Gasto.empresa_id == empresa_id)
        existente = q_dup.first()

    if existente:
        resultado["estado"] = "duplicado"
        resultado["ya_registrado"] = True
        resultado["mensaje_duplicado"] = (
            f"Este gasto ya está registrado: {existente.numero_comprobante} "
            f"del {existente.fecha.strftime('%d/%m/%Y')}"
        )
    else:
        resultado["ya_registrado"] = False
        resultado["mensaje_duplicado"] = None
        if not resultado["tipo_comprobante"]:
            resultado["estado"] = "no_valido"
            resultado["error"] = "No se reconoce como una factura o boleta de proveedor (SUNAT)"
        else:
            campos_criticos_faltantes = [c for c in resultado["campos_no_detectados"] if c not in ("categoria", "area")]
            resultado["estado"] = "revisar" if campos_criticos_faltantes else "listo"

    return resultado


@router.post("/importar-pdf")
async def importar_pdf_gasto(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(400, "El archivo debe ser un PDF")

    contenido = await file.read()
    if len(contenido) > TAMANO_MAXIMO_PDF_GASTO:
        raise HTTPException(400, "El archivo supera el tamaño máximo de 10 MB")

    ruc_propio = _ruc_empresa_propia(db, empresa_id)
    resultado = _procesar_pdf_individual_gasto(contenido, file.filename, db, ruc_propio, empresa_id)
    if resultado.get("no_valido_ilegible"):
        raise HTTPException(422, resultado["error"])
    return resultado


@router.post("/importar-zip")
async def importar_zip_gasto(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    if not (file.filename or "").lower().endswith(".zip"):
        raise HTTPException(400, "El archivo debe ser un ZIP")

    contenido = await file.read()
    if len(contenido) > TAMANO_MAXIMO_ZIP_GASTO:
        raise HTTPException(400, "El archivo ZIP supera el tamaño máximo de 50 MB")

    try:
        zf = zipfile.ZipFile(io.BytesIO(contenido))
    except zipfile.BadZipFile:
        raise HTTPException(400, "El archivo no es un ZIP válido")

    pdf_names = [n for n in zf.namelist() if n.lower().endswith(".pdf") and not n.endswith("/")]
    if not pdf_names:
        raise HTTPException(400, "El ZIP no contiene archivos PDF")
    if len(pdf_names) > MAX_PDFS_POR_ZIP_GASTO:
        raise HTTPException(400, f"El ZIP contiene {len(pdf_names)} PDFs; el máximo permitido es {MAX_PDFS_POR_ZIP_GASTO}")

    ruc_propio = _ruc_empresa_propia(db, empresa_id)
    resultados = []
    for pdf_name in pdf_names:
        try:
            with zf.open(pdf_name) as pf:
                contenido_pdf = pf.read()
        except RuntimeError as e:
            if "password" in str(e).lower():
                raise HTTPException(
                    400,
                    "El archivo ZIP está protegido con contraseña, por favor descomprímelo primero",
                )
            raise
        nombre_corto = pdf_name.rsplit("/", 1)[-1]
        resultado = _procesar_pdf_individual_gasto(contenido_pdf, nombre_corto, db, ruc_propio, empresa_id)
        resultado["archivo"] = nombre_corto

        TMP_ZIP_DIR_GASTOS.mkdir(parents=True, exist_ok=True)
        nombre_temp = f"{uuid.uuid4().hex}.pdf"
        with open(TMP_ZIP_DIR_GASTOS / nombre_temp, "wb") as f:
            f.write(contenido_pdf)
        resultado["archivo_temp"] = nombre_temp

        resultados.append(resultado)

    return {"total": len(pdf_names), "resultados": resultados}


class ImportarZipConfirmarRequestGasto(BaseModel):
    gastos: List[GastoCreate]


@router.post("/importar-zip/confirmar")
def importar_zip_confirmar_gasto(data: ImportarZipConfirmarRequestGasto, http_request: Request,
                                  db: Session = Depends(get_db),
                                  usuario: Usuario = Depends(get_current_usuario),
                                  empresa_id: Optional[int] = Depends(get_empresa_id)):
    resultados = []
    for item in data.gastos:
        try:
            r = _crear_gasto(item, http_request, db, usuario, empresa_id)
            resultados.append({
                "numero_comprobante": item.numero_comprobante, "exito": True, "id": r["id"],
            })
        except HTTPException as e:
            resultados.append({
                "numero_comprobante": item.numero_comprobante, "exito": False, "error": e.detail,
            })

    exitosos = sum(1 for r in resultados if r["exito"])
    return {"total": len(resultados), "exitosos": exitosos, "resultados": resultados}


# ── Cuentas por Pagar ────────────────────────────────────────────────────────────

@router.get("/cuentas-por-pagar")
def listar_cpp(
    semaforo:  str            = "",
    search:    str            = "",
    categoria: str            = "",
    area:      str            = "",
    estado:    str            = "",
    desde:     Optional[date] = None,
    hasta:     Optional[date] = None,
    page:      int            = 1,
    per_page:  int            = 20,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.tipo_comprobante != "Anticipo de Proveedor")
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    if search:
        like = f"%{search}%"
        q = q.filter(
            Gasto.descripcion.ilike(like) |
            Gasto.proveedor.ilike(like) |
            Gasto.numero_comprobante.ilike(like) |
            Gasto.numero_documento.ilike(like)
        )
    if categoria: q = q.filter(Gasto.categoria == categoria)
    if area:      q = q.filter(Gasto.area == area)
    if estado:    q = q.filter(Gasto.estado_pago == estado)
    if desde:     q = q.filter(Gasto.fecha >= desde)
    if hasta:     q = q.filter(Gasto.fecha <= hasta)
    all_rows = q.order_by(Gasto.fecha.desc(), Gasto.id.desc()).all()
    if semaforo:
        all_rows = [g for g in all_rows if _semaforo_gasto(g) == semaforo]
    total     = len(all_rows)
    offset    = (page - 1) * per_page
    page_rows = all_rows[offset : offset + per_page]
    return {"total": total, "page": page, "per_page": per_page, "data": [_serialize_cpp(g) for g in page_rows]}


@router.get("/cuentas-por-pagar/resumen")
def resumen_cpp(
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.tipo_comprobante != "Anticipo de Proveedor")
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    all_gastos    = q.all()
    total_pend    = al_dia = por_vencer_15 = vencido_mas_15 = 0.0
    for g in all_gastos:
        if g.estado_pago == "Pagado":
            continue
        saldo = float(g.saldo_pendiente) if g.saldo_pendiente is not None else float(g.monto or 0)
        total_pend += saldo
        s = _semaforo_gasto(g)
        if s == "verde":
            al_dia += saldo
        elif s == "amarillo":
            por_vencer_15 += saldo
        elif s == "rojo":
            vencido_mas_15 += saldo
    return {
        "total_pendiente": round(total_pend, 2),
        "al_dia":          round(al_dia, 2),
        "por_vencer_15":   round(por_vencer_15, 2),
        "vencido_mas_15":  round(vencido_mas_15, 2),
    }


@router.get("/cuentas-por-pagar/exportar")
def exportar_cpp(
    desde:     Optional[date] = None,
    hasta:     Optional[date] = None,
    proveedor: str            = "",
    estado:    str            = "",
    categoria: str            = "",
    area:      str            = "",
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(500, "Instale openpyxl: pip install openpyxl")

    q = db.query(Gasto).filter(Gasto.tipo_comprobante != "Anticipo de Proveedor")
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    if desde:     q = q.filter(Gasto.fecha >= desde)
    if hasta:     q = q.filter(Gasto.fecha <= hasta)
    if proveedor: q = q.filter(Gasto.proveedor.ilike(f"%{proveedor}%"))
    if estado:    q = q.filter(Gasto.estado_pago == estado)
    if categoria: q = q.filter(Gasto.categoria == categoria)
    if area:      q = q.filter(Gasto.area == area)
    rows  = q.order_by(Gasto.fecha.desc()).all()
    hoy_d = date.today()

    wb = Workbook()
    ws = wb.active
    ws.title = "Cuentas por Pagar"

    hdr_font   = Font(bold=True, color="FFFFFF")
    hdr_fill   = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    hdr_align  = Alignment(horizontal="center", vertical="center")
    alt_fill   = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    title_font = Font(bold=True, size=12, color="1E40AF")

    HEADERS = ["Fecha", "Categoría", "Área", "N° Comprobante", "Proveedor", "Descripción",
               "RUC/Doc", "Tipo Comprobante", "Moneda", "Monto Original", "T/C",
               "Base Imponible (S/)", "IGV 18% (S/)", "Monto Total (S/)", "Saldo Pendiente (S/)",
               "Estado", "F. Vencimiento", "Días de Mora"]
    WIDTHS  = [12, 22, 16, 16, 28, 30, 16, 20, 10, 14, 10, 18, 14, 14, 22, 14, 16, 14]

    filtros_label = []
    if categoria: filtros_label.append(f"Categoría: {categoria}")
    if area:      filtros_label.append(f"Área: {area}")
    titulo = "Reporte Cuentas por Pagar"
    if filtros_label:
        titulo += " — " + " | ".join(filtros_label)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title_cell = ws.cell(row=1, column=1, value=titulo)
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    for ci, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=2, column=ci, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = hdr_align
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[2].height = 20

    for ri, g in enumerate(rows, 3):
        saldo     = float(g.saldo_pendiente) if g.saldo_pendiente is not None else float(g.monto or 0)
        dias_mora = 0
        if g.fecha_vencimiento and g.estado_pago != "Pagado":
            dias_mora = max(0, (hoy_d - g.fecha_vencimiento).days)
        num_doc   = f"{g.tipo_documento}: {g.numero_documento}" if g.tipo_documento and g.numero_documento else (g.numero_documento or "")
        moneda_g  = g.moneda or "PEN"
        monto_s   = round(float(g.monto_soles or g.monto), 2)
        monto_o   = round(float(g.monto_original), 2) if g.monto_original is not None else monto_s
        tc_g      = round(float(g.tipo_cambio), 4) if g.tipo_cambio is not None else None
        ws.append([
            str(g.fecha),
            g.categoria or "—",
            g.area or "—",
            g.numero_comprobante or "—",
            g.proveedor or "—",
            g.descripcion or "—",
            num_doc or "—",
            g.tipo_comprobante or "—",
            moneda_g,
            monto_o,
            tc_g,
            round(float(g.base_imponible), 2) if g.base_imponible is not None else None,
            round(float(g.igv), 2) if g.igv is not None else None,
            monto_s,
            round(saldo, 2),
            g.estado_pago or "Pendiente",
            str(g.fecha_vencimiento) if g.fecha_vencimiento else "—",
            dias_mora,
        ])
        if ri % 2 == 1:
            for ci2 in range(1, len(HEADERS) + 1):
                ws.cell(row=ri, column=ci2).fill = alt_fill

    ws.freeze_panes = "A3"
    output = io.BytesIO()
    wb.save(output); output.seek(0)
    fname = f"Cuentas_Por_Pagar_{date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


# ── Pagos Tributarios ─────────────────────────────────────────────────────────────

@router.get("/pagos-tributarios")
def listar_pagos_tributarios(
    concepto:     str            = "",
    periodo_mes:  Optional[int]  = None,
    periodo_anio: Optional[int]  = None,
    estado:       str            = "",
    search:       str            = "",
    page:         int            = 1,
    per_page:     int            = 20,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.categoria == CATEGORIA_NO_AFECTA_UTILIDAD)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    if concepto:     q = q.filter(Gasto.descripcion == concepto)
    if periodo_mes:  q = q.filter(Gasto.periodo_mes == periodo_mes)
    if periodo_anio: q = q.filter(Gasto.periodo_anio == periodo_anio)
    if estado:       q = q.filter(Gasto.estado_pago == estado)
    if search:
        like = f"%{search}%"
        q = q.filter(Gasto.descripcion.ilike(like) | Gasto.observaciones.ilike(like))
    all_rows = q.order_by(Gasto.fecha_vencimiento.asc(), Gasto.id.desc()).all()
    total     = len(all_rows)
    offset    = (page - 1) * per_page
    page_rows = all_rows[offset : offset + per_page]
    return {"total": total, "page": page, "per_page": per_page, "data": [_serialize_tributario(db, g) for g in page_rows]}


@router.get("/pagos-tributarios/resumen-kpis")
def resumen_pagos_tributarios(
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    hoy = date.today()

    q_pagado_mes = (
        db.query(func.sum(PagoGasto.monto_pagado))
        .join(Gasto, Gasto.id == PagoGasto.gasto_id)
        .filter(
            Gasto.categoria == CATEGORIA_NO_AFECTA_UTILIDAD,
            extract("month", PagoGasto.fecha_pago) == hoy.month,
            extract("year",  PagoGasto.fecha_pago) == hoy.year,
        )
    )
    if empresa_id is not None:
        q_pagado_mes = q_pagado_mes.filter(Gasto.empresa_id == empresa_id)
    pagado_mes = float(q_pagado_mes.scalar() or 0)

    q_pagado_anio = (
        db.query(func.sum(PagoGasto.monto_pagado))
        .join(Gasto, Gasto.id == PagoGasto.gasto_id)
        .filter(
            Gasto.categoria == CATEGORIA_NO_AFECTA_UTILIDAD,
            extract("year", PagoGasto.fecha_pago) == hoy.year,
        )
    )
    if empresa_id is not None:
        q_pagado_anio = q_pagado_anio.filter(Gasto.empresa_id == empresa_id)
    pagado_anio = float(q_pagado_anio.scalar() or 0)

    q_pendiente = db.query(func.sum(Gasto.saldo_pendiente)).filter(
        Gasto.categoria == CATEGORIA_NO_AFECTA_UTILIDAD,
        Gasto.estado_pago != "Pagado",
    )
    if empresa_id is not None:
        q_pendiente = q_pendiente.filter(Gasto.empresa_id == empresa_id)
    pendiente = float(q_pendiente.scalar() or 0)

    return {
        "pagado_mes":  round(pagado_mes, 2),
        "pendiente":   round(pendiente, 2),
        "pagado_anio": round(pagado_anio, 2),
    }


@router.get("/pagos-tributarios/exportar")
def exportar_pagos_tributarios(
    periodo_mes:  Optional[int] = None,
    periodo_anio: Optional[int] = None,
    concepto:     str           = "",
    estado:       str           = "",
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(500, "Instale openpyxl: pip install openpyxl")

    q = db.query(Gasto).filter(Gasto.categoria == CATEGORIA_NO_AFECTA_UTILIDAD)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    if periodo_mes:  q = q.filter(Gasto.periodo_mes == periodo_mes)
    if periodo_anio: q = q.filter(Gasto.periodo_anio == periodo_anio)
    if concepto:     q = q.filter(Gasto.descripcion == concepto)
    if estado:       q = q.filter(Gasto.estado_pago == estado)
    rows = q.order_by(Gasto.fecha_vencimiento.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Pagos Tributarios"

    hdr_font  = Font(bold=True, color="FFFFFF")
    hdr_fill  = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center")
    alt_fill  = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    HEADERS = ["Concepto", "Período", "Monto (S/)", "Fecha Límite", "Fecha Pago", "Estado"]
    WIDTHS  = [24, 16, 14, 16, 16, 14]

    for ci, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = hdr_align
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 20

    for ri, g in enumerate(rows, 2):
        d = _serialize_tributario(db, g)
        ws.append([
            d["concepto"],
            d["periodo_label"] or "—",
            d["monto"],
            d["fecha_limite"] or "—",
            d["fecha_pago"] or "—",
            d["estado"],
        ])
        if ri % 2 == 0:
            for ci2 in range(1, len(HEADERS) + 1):
                ws.cell(row=ri, column=ci2).fill = alt_fill

    ws.freeze_panes = "A2"
    output = io.BytesIO()
    wb.save(output); output.seek(0)
    fname = f"Pagos_Tributarios_{date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )

# ── Editar / Eliminar pagos ───────────────────────────────────────────────────────

@router.put("/pagos/{pago_id}")
def editar_pago_gasto(
    pago_id: int,
    data: PagoGastoUpdate,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    pago = db.query(PagoGasto).filter(PagoGasto.id == pago_id).first()
    if not pago:
        raise HTTPException(404, "Pago no encontrado")
    q_g = db.query(Gasto).filter(Gasto.id == pago.gasto_id)
    if empresa_id is not None:
        q_g = q_g.filter(Gasto.empresa_id == empresa_id)
    g = q_g.first()
    if not g:
        raise HTTPException(404, "Gasto no encontrado")
    if data.monto_pagado <= 0:
        raise HTTPException(400, "El monto debe ser mayor a 0")
    pago.monto_pagado  = round(data.monto_pagado, 2)
    pago.fecha_pago    = data.fecha_pago
    pago.metodo_pago   = data.metodo_pago
    pago.banco         = data.banco or None
    pago.numero_cuenta = data.numero_cuenta or None
    pago.numero_cheque = data.numero_cheque or None
    pago.numero_operacion = data.numero_operacion or None
    db.flush()
    _recalcular_pago_gasto(db, g)
    db.commit()
    db.refresh(g)
    return {"mensaje": "Pago actualizado", "saldo_pendiente": g.saldo_pendiente, "estado_pago": g.estado_pago}


@router.delete("/pagos/{pago_id}")
def eliminar_pago_gasto(
    pago_id: int,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    pago = db.query(PagoGasto).filter(PagoGasto.id == pago_id).first()
    if not pago:
        raise HTTPException(404, "Pago no encontrado")

    if pago.tipo == "pago_cuota_prestamo":
        cuota = db.query(CuotaPrestamo).filter(CuotaPrestamo.id == pago.referencia_id).first()
        if cuota:
            prestamo = db.query(Prestamo).filter(Prestamo.id == cuota.prestamo_id).first()
            if prestamo:
                prestamo.monto_pendiente = round(float(prestamo.monto_pendiente) + float(cuota.amortizacion), 2)
                if prestamo.estado == "pagado":
                    prestamo.estado = "activo"
            movimiento_caja_id_anterior = cuota.movimiento_caja_id
            cuota.estado             = "pendiente"
            cuota.fecha_pago_real    = None
            cuota.metodo_pago        = None
            cuota.cuenta_bancaria_id = None
            cuota.movimiento_caja_id = None
            db.flush()
            if movimiento_caja_id_anterior:
                movimiento = db.query(MovimientoCaja).filter(MovimientoCaja.id == movimiento_caja_id_anterior).first()
                if movimiento:
                    db.delete(movimiento)
            db.query(Gasto).filter(Gasto.cuota_prestamo_id == cuota.id).delete()
            limpiar_movimiento_sistema_por_pago(db, "prestamo", cuota.id)
        limpiar_movimiento_sistema_por_pago(db, "gasto", pago.id)
        db.delete(pago)
        db.commit()
        return {"mensaje": "Pago de cuota eliminado — la cuota volvió a estado pendiente"}

    q_g = db.query(Gasto).filter(Gasto.id == pago.gasto_id)
    if empresa_id is not None:
        q_g = q_g.filter(Gasto.empresa_id == empresa_id)
    g = q_g.first()
    if not g:
        raise HTTPException(404, "Gasto no encontrado")
    limpiar_movimiento_sistema_por_pago(db, "gasto", pago.id)
    db.delete(pago)
    db.flush()
    _recalcular_pago_gasto(db, g)
    db.commit()
    db.refresh(g)
    return {"mensaje": "Pago eliminado", "saldo_pendiente": g.saldo_pendiente, "estado_pago": g.estado_pago}


# ── Lista de Pagos ────────────────────────────────────────────────────────────────

@router.get("/pagos")
def listar_pagos(
    desde:         Optional[date] = None,
    hasta:         Optional[date] = None,
    banco:         str            = "",
    numero_cuenta: str            = "",
    metodo_pago:   str            = "",
    page:          int            = 1,
    per_page:      int            = 20,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    filas = []

    # PagoGasto: usa outerjoin con Gasto para filtrar por empresa_id,
    # manteniendo los pago_cuota_prestamo (gasto_id=None) que no tienen Gasto.
    q = db.query(PagoGasto)
    if empresa_id is not None:
        from sqlalchemy.orm import aliased
        GastoAlias = aliased(Gasto)
        q = (
            db.query(PagoGasto)
            .outerjoin(GastoAlias, PagoGasto.gasto_id == GastoAlias.id)
            .filter(
                (PagoGasto.gasto_id == None) | (GastoAlias.empresa_id == empresa_id)
            )
        )
    if desde:         q = q.filter(PagoGasto.fecha_pago >= desde)
    if hasta:         q = q.filter(PagoGasto.fecha_pago <= hasta)
    if banco:         q = q.filter(PagoGasto.banco == banco)
    if numero_cuenta: q = q.filter(PagoGasto.numero_cuenta == numero_cuenta)
    if metodo_pago:   q = q.filter(PagoGasto.metodo_pago == metodo_pago)
    for p in q.all():
        if p.tipo == "pago_cuota_prestamo":
            cuota    = db.query(CuotaPrestamo).filter(CuotaPrestamo.id == p.referencia_id).first()
            prestamo = db.query(Prestamo).filter(Prestamo.id == cuota.prestamo_id).first() if cuota else None
            if cuota and prestamo:
                if prestamo.tipo in ("recibido_banco", "recibido_tercero"):
                    descripcion = f"Pago cuota préstamo — {prestamo.nombre_tercero} — Cuota N° {cuota.numero_cuota}"
                else:
                    descripcion = f"Cuota préstamo otorgado — {prestamo.nombre_tercero} — Cuota N° {cuota.numero_cuota}"
                numero_comprobante = f"PREST-{prestamo.id}-C{cuota.numero_cuota}"
                proveedor = prestamo.nombre_tercero
            else:
                descripcion = "Cuota de préstamo (registro no encontrado)"
                numero_comprobante = None
                proveedor = None
            filas.append({
                "tipo":               "cuota_prestamo",
                "id":                 p.id,
                "gasto_id":           None,
                "lote_id":            None,
                "orden_id":           None,
                "fecha_pago":         p.fecha_pago,
                "numero_comprobante": numero_comprobante,
                "proveedor":          proveedor,
                "descripcion":        descripcion,
                "monto_pagado":       round(float(p.monto_pagado), 2),
                "banco":              p.banco,
                "numero_cuenta":      p.numero_cuenta,
                "metodo_pago":        p.metodo_pago,
                "numero_cheque":      None,
                "numero_operacion":   None,
                "cantidad_facturas":  None,
            })
            continue
        filas.append({
            "tipo":               "gasto",
            "id":                 p.id,
            "gasto_id":           p.gasto_id,
            "lote_id":            None,
            "orden_id":           None,
            "fecha_pago":         p.fecha_pago,
            "numero_comprobante": p.gasto.numero_comprobante if p.gasto else None,
            "proveedor":          p.gasto.proveedor if p.gasto else None,
            "descripcion":        p.gasto.descripcion if p.gasto else None,
            "monto_pagado":       round(float(p.monto_pagado), 2),
            "banco":              p.banco,
            "numero_cuenta":      p.numero_cuenta,
            "metodo_pago":        p.metodo_pago,
            "numero_cheque":      p.numero_cheque,
            "numero_operacion":   p.numero_operacion,
            "cantidad_facturas":  None,
        })

    ql = db.query(LoteDetraccion).filter(
        LoteDetraccion.estado == "pagado", LoteDetraccion.fecha_pago.isnot(None),
    )
    if empresa_id is not None:
        ql = ql.filter(LoteDetraccion.empresa_id == empresa_id)
    if desde:         ql = ql.filter(LoteDetraccion.fecha_pago >= desde)
    if hasta:         ql = ql.filter(LoteDetraccion.fecha_pago <= hasta)
    if banco:         ql = ql.filter(LoteDetraccion.banco == banco)
    if numero_cuenta: ql = ql.filter(LoteDetraccion.numero_cuenta == numero_cuenta)
    if metodo_pago:   ql = ql.filter(LoteDetraccion.metodo_pago == metodo_pago)
    for l in ql.all():
        filas.append({
            "tipo":               "lote_detraccion",
            "id":                 l.id,
            "gasto_id":           None,
            "lote_id":            l.id,
            "orden_id":           None,
            "fecha_pago":         l.fecha_pago,
            "numero_comprobante": f"LOTE-{l.numero_lote}",
            "proveedor":          None,
            "descripcion":        "PAGO DE DETRACCIONES",
            "monto_pagado":       round(float(l.importe_total or 0), 2),
            "banco":              l.banco,
            "numero_cuenta":      l.numero_cuenta,
            "metodo_pago":        l.metodo_pago,
            "numero_cheque":      None,
            "numero_operacion":   l.numero_operacion,
            "cantidad_facturas":  len(l.detalles),
        })

    qo = db.query(OrdenPago)
    if empresa_id is not None:
        qo = qo.filter(OrdenPago.empresa_id == empresa_id)
    if desde:         qo = qo.filter(OrdenPago.fecha_pago >= desde)
    if hasta:         qo = qo.filter(OrdenPago.fecha_pago <= hasta)
    if banco:         qo = qo.filter(OrdenPago.banco == banco)
    if numero_cuenta: qo = qo.filter(OrdenPago.numero_cuenta == numero_cuenta)
    if metodo_pago:   qo = qo.filter(OrdenPago.metodo_pago == metodo_pago)
    for o in qo.all():
        filas.append({
            "tipo":               "orden_pago",
            "id":                 o.id,
            "gasto_id":           None,
            "lote_id":            None,
            "orden_id":           o.id,
            "fecha_pago":         o.fecha_pago,
            "numero_comprobante": o.numero_orden,
            "proveedor":          None,
            "descripcion":        "ORDEN DE PAGO",
            "monto_pagado":       round(float(o.monto_total or 0), 2),
            "banco":              o.banco,
            "numero_cuenta":      o.numero_cuenta,
            "metodo_pago":        o.metodo_pago,
            "numero_cheque":      o.numero_cheque,
            "numero_operacion":   o.numero_operacion,
            "cantidad_facturas":  len(o.detalles),
        })

    qg = db.query(PagoGarantia, Garantia).join(
        Garantia, PagoGarantia.garantia_id == Garantia.id
    ).filter(PagoGarantia.tipo == "devolucion")
    if empresa_id is not None:
        qg = qg.filter(Garantia.empresa_id == empresa_id)
    if desde:       qg = qg.filter(PagoGarantia.fecha >= desde)
    if hasta:       qg = qg.filter(PagoGarantia.fecha <= hasta)
    if metodo_pago: qg = qg.filter(PagoGarantia.metodo_pago == metodo_pago)
    if banco or numero_cuenta:
        qg = qg.join(CuentaBancaria, PagoGarantia.cuenta_bancaria_id == CuentaBancaria.id)
        if banco:         qg = qg.filter(CuentaBancaria.banco == banco)
        if numero_cuenta: qg = qg.filter(CuentaBancaria.numero_cuenta == numero_cuenta)
    for pg, gar in qg.all():
        cuenta = pg.cuenta_bancaria
        filas.append({
            "tipo":               "devolucion_garantia",
            "id":                 pg.id,
            "gasto_id":           None,
            "lote_id":            None,
            "orden_id":           None,
            "fecha_pago":         pg.fecha,
            "numero_comprobante": pg.numero_documento_generado,
            "proveedor":          gar.cliente_nombre,
            "descripcion":        f"Devolución garantía — {gar.cliente_nombre}",
            "monto_pagado":       round(float(pg.monto), 2),
            "banco":              cuenta.banco if cuenta else None,
            "numero_cuenta":      cuenta.numero_cuenta if cuenta else None,
            "metodo_pago":        pg.metodo_pago,
            "numero_cheque":      None,
            "numero_operacion":   None,
            "cantidad_facturas":  None,
        })

    filas.sort(key=lambda f: f["fecha_pago"] or date.min, reverse=True)
    total     = len(filas)
    offset    = (page - 1) * per_page
    page_rows = filas[offset : offset + per_page]
    for f in page_rows:
        f["fecha_pago"] = str(f["fecha_pago"]) if f["fecha_pago"] else None
    return {"total": total, "page": page, "per_page": per_page, "data": page_rows}


# ── Obtener gasto ────────────────────────────────────────────────────────────────

@router.get("/{gasto_id}")
def obtener_gasto(
    gasto_id: int,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.id == gasto_id)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    g = q.first()
    if not g:
        raise HTTPException(404, "Gasto no encontrado")
    return _serialize(g)


# ── Comprobante adjunto ──────────────────────────────────────────────────────────

@router.post("/{gasto_id}/comprobante")
async def subir_comprobante_gasto(
    gasto_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.id == gasto_id)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    g = q.first()
    if not g:
        raise HTTPException(404, "Gasto no encontrado")

    if g.tipo_comprobante not in TIPOS_CON_PDF_GASTO:
        return {"mensaje": "Este tipo de comprobante no conserva el PDF original", "adjuntado": False}

    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXT_GASTOS:
        raise HTTPException(400, "Solo se aceptan PDF, JPG y PNG")

    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "Archivo mayor a 10 MB")

    if g.comprobante_path:
        Path(g.comprobante_path).unlink(missing_ok=True)

    upload_dir = UPLOAD_BASE_GASTOS / str(gasto_id)
    upload_dir.mkdir(parents=True, exist_ok=True)
    file_path = upload_dir / file.filename

    with open(file_path, "wb") as f:
        f.write(content)

    g.comprobante_path   = str(file_path)
    g.comprobante_nombre = file.filename
    db.commit()

    return {"mensaje": "Comprobante subido", "nombre": file.filename, "adjuntado": True}


@router.get("/{gasto_id}/comprobante")
def ver_comprobante_gasto(
    gasto_id: int,
    download: bool = False,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.id == gasto_id)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    g = q.first()
    if not g or not g.comprobante_path:
        raise HTTPException(404, "Comprobante no encontrado")
    if not Path(g.comprobante_path).exists():
        raise HTTPException(404, "Archivo no encontrado en disco")
    headers = {}
    if download:
        headers["Content-Disposition"] = f'attachment; filename="{g.comprobante_nombre}"'
    else:
        headers["Content-Disposition"] = f'inline; filename="{g.comprobante_nombre}"'
    return FileResponse(g.comprobante_path, headers=headers)


@router.delete("/{gasto_id}/comprobante")
def eliminar_comprobante_gasto(
    gasto_id: int,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.id == gasto_id)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    g = q.first()
    if not g:
        raise HTTPException(404, "Gasto no encontrado")
    if g.comprobante_path:
        Path(g.comprobante_path).unlink(missing_ok=True)
    g.comprobante_path   = None
    g.comprobante_nombre = None
    db.commit()
    return {"mensaje": "Comprobante eliminado"}


# ── Imprimir ──────────────────────────────────────────────────────────────────────

@router.get("/{gasto_id}/imprimir")
def imprimir_gasto(
    gasto_id: int,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_current_usuario),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.id == gasto_id)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    g = q.first()
    if not g:
        raise HTTPException(404, "Gasto no encontrado")
    s = _serialize(g)

    proveedor_direccion = None
    if g.proveedor_id:
        prov = db.query(Proveedor).filter(Proveedor.id == g.proveedor_id).first()
        proveedor_direccion = prov.direccion if prov else None

    monto_total = s["monto_soles"]
    tiene_igv = s["base_imponible"] is not None and s["igv"] is not None

    items = [{
        "cantidad":    1,
        "unidad":      "SERVICIO",
        "codigo":      "",
        "descripcion": s["descripcion"] or s["categoria"] or "Gasto",
        "v_unit":      s["base_imponible"] if tiene_igv else monto_total,
        "igv":         s["igv"] if tiene_igv else 0,
        "p_unit":      monto_total,
        "total":       monto_total,
    }]

    return construir_payload_impresion(
        db, usuario,
        tipo_documento=s["tipo_comprobante"] or "Recibo Interno",
        numero_documento=s["numero_comprobante"],
        fecha_emision=g.fecha,
        fecha_vencimiento=g.fecha_vencimiento,
        moneda=s["moneda"],
        cliente_nombre=s["proveedor"] or "—",
        cliente_ruc=s["numero_documento"],
        cliente_direccion=proveedor_direccion,
        cliente_label="PROVEEDOR / BENEFICIARIO",
        items=items,
        base_imponible=s["base_imponible"] if tiene_igv else 0,
        igv=s["igv"] if tiene_igv else 0,
        total=monto_total,
    )


# ── Actualizar gasto ─────────────────────────────────────────────────────────────

@router.put("/{gasto_id}")
def actualizar_gasto(
    gasto_id: int,
    data: GastoUpdate,
    http_request: Request,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_current_usuario),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.id == gasto_id)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    g = q.first()
    if not g:
        raise HTTPException(404, "Gasto no encontrado")
    if data.monto <= 0:
        raise HTTPException(400, "El monto debe ser mayor a 0")

    numero = data.numero_comprobante
    if data.tipo_comprobante == "Recibo Interno" and not numero:
        numero = _proximo_ri(db, empresa_id)
    elif data.tipo_comprobante == "Gastos Bancarios" and not numero:
        numero = _proximo_gb(db, empresa_id)
    elif data.tipo_comprobante == "Anticipo de Proveedor" and not numero:
        numero = _proximo_ap(db, empresa_id)

    es_anticipo = data.tipo_comprobante == "Anticipo de Proveedor"

    moneda_u   = data.moneda or "PEN"
    tc_u       = round(data.tipo_cambio, 4) if data.tipo_cambio else None
    monto_r_u  = round(data.monto, 2)
    monto_orig_u = monto_r_u if moneda_u == "USD" else None
    m_soles_u    = round(monto_r_u * tc_u, 2) if moneda_u == "USD" and tc_u else monto_r_u
    base_u, igv_u = _calcular_impuestos(m_soles_u, data.tipo_comprobante)

    g.fecha              = data.fecha
    g.categoria          = data.categoria
    g.descripcion        = data.descripcion.strip()
    g.monto              = m_soles_u
    g.area               = data.area
    g.es_recurrente      = data.es_recurrente
    g.tipo_comprobante   = data.tipo_comprobante or None
    g.numero_comprobante = numero or None
    g.proveedor          = data.proveedor or None
    g.tipo_documento     = data.tipo_documento or None
    g.numero_documento   = data.numero_documento or None
    g.fecha_vencimiento  = data.fecha_vencimiento or None
    g.base_imponible     = base_u
    g.igv                = igv_u
    g.moneda             = moneda_u
    g.tipo_cambio        = tc_u
    g.monto_original     = monto_orig_u
    g.monto_soles        = m_soles_u
    g.orden_id           = data.orden_id
    g.afecta_utilidad    = _afecta_utilidad(data.categoria)
    g.periodo_mes        = data.periodo_mes
    g.periodo_anio       = data.periodo_anio
    g.observaciones      = data.observaciones or None

    tiene_det, tasa_det, monto_det, monto_neto, fecha_lim_det = _calcular_detraccion_gasto(
        data.tipo_comprobante, data.tiene_detraccion, data.tasa_detraccion, m_soles_u,
        data.fecha, data.fecha_limite_detraccion,
    )
    g.tiene_detraccion        = tiene_det
    g.tasa_detraccion         = tasa_det
    g.monto_detraccion        = monto_det
    g.monto_neto_pagar        = monto_neto
    g.concepto_detraccion     = data.concepto_detraccion if tiene_det else None
    g.ruc_cuenta_detraccion   = data.ruc_cuenta_detraccion if tiene_det else None
    g.fecha_limite_detraccion = fecha_lim_det
    g.codigo_detraccion       = data.codigo_detraccion if data.tipo_comprobante in TIPOS_CON_CODIGO_DETRACCION else None

    if data.es_recurrente:
        g.recurrente_activo = True
    if es_anticipo:
        g.saldo_pendiente = None
        g.estado_pago     = None
    else:
        _recalcular_pago_gasto(db, g)

    g.proveedor_id = _upsert_proveedor(db, data.tipo_documento, data.numero_documento, data.proveedor, empresa_id)
    g.modificado_por = usuario.nombre
    g.modificado_en  = datetime.utcnow()
    db.commit()
    db.refresh(g)

    registrar_log(
        db, usuario.id, usuario.nombre, "gastos", "Editó gasto",
        f"Editó gasto {g.numero_comprobante or g.descripcion} - S/ {float(g.monto or 0):,.2f}",
        ip_de(http_request),
    )

    return _serialize(g)


# ── Eliminar gasto ───────────────────────────────────────────────────────────────

class EliminarGastosMasivoRequest(BaseModel):
    ids: List[int]


@router.delete("/eliminar-masivo")
def eliminar_gastos_masivo(
    data: EliminarGastosMasivoRequest,
    http_request: Request,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_current_usuario),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    resultados = []
    for gasto_id in data.ids:
        q = db.query(Gasto).filter(Gasto.id == gasto_id)
        if empresa_id is not None:
            q = q.filter(Gasto.empresa_id == empresa_id)
        g = q.first()
        if not g:
            resultados.append({"id": gasto_id, "exito": False, "error": "Gasto no encontrado"})
            continue

        try:
            db.query(LoteDetraccionDetalle).filter(
                LoteDetraccionDetalle.gasto_id == gasto_id
            ).delete()

            db.query(OrdenPagoDetalle).filter(
                OrdenPagoDetalle.gasto_id == gasto_id
            ).delete()

            db.query(PagoGasto).filter(
                PagoGasto.gasto_id == gasto_id
            ).delete()

            if g.es_recurrente:
                for inst in db.query(Gasto).filter(Gasto.recurrente_padre_id == gasto_id).all():
                    db.delete(inst)

            identificador = g.numero_comprobante or g.descripcion
            db.delete(g)
            db.commit()

            registrar_log(
                db, usuario.id, usuario.nombre, "gastos", "Eliminó gasto",
                f"Eliminó gasto {identificador} (eliminación masiva)", ip_de(http_request),
            )
            resultados.append({"id": gasto_id, "exito": True, "identificador": identificador})
        except IntegrityError as e:
            db.rollback()
            resultados.append({"id": gasto_id, "exito": False, "error": f"No se puede eliminar: {str(e)}"})

    exitosos = sum(1 for r in resultados if r["exito"])
    return {"total": len(data.ids), "exitosos": exitosos, "resultados": resultados}


@router.delete("/{gasto_id}")
def eliminar_gasto(
    gasto_id: int,
    http_request: Request,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_current_usuario),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.id == gasto_id)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    g = q.first()
    if not g:
        raise HTTPException(404, "Gasto no encontrado")

    db.query(LoteDetraccionDetalle).filter(
        LoteDetraccionDetalle.gasto_id == gasto_id
    ).delete()

    db.query(OrdenPagoDetalle).filter(
        OrdenPagoDetalle.gasto_id == gasto_id
    ).delete()

    if g.es_recurrente:
        for inst in db.query(Gasto).filter(Gasto.recurrente_padre_id == gasto_id).all():
            db.delete(inst)
    identificador = g.numero_comprobante or g.descripcion
    try:
        db.delete(g)
        db.commit()
    except IntegrityError as e:
        db.rollback()
        raise HTTPException(400, f"No se puede eliminar: {str(e)}")

    registrar_log(
        db, usuario.id, usuario.nombre, "gastos", "Eliminó gasto",
        f"Eliminó gasto {identificador}", ip_de(http_request),
    )

    return {"mensaje": "Gasto eliminado correctamente"}


# ── Toggle recurrente ────────────────────────────────────────────────────────────

@router.post("/{gasto_id}/toggle-recurrente")
def toggle_recurrente(
    gasto_id: int,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.id == gasto_id)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    g = q.first()
    if not g:
        raise HTTPException(404, "Gasto no encontrado")
    if not g.es_recurrente:
        raise HTTPException(400, "Este gasto no está marcado como recurrente")
    g.recurrente_activo = not g.recurrente_activo
    db.commit()
    db.refresh(g)
    return {"mensaje": f"Recurrente {'activado' if g.recurrente_activo else 'desactivado'}", "recurrente_activo": g.recurrente_activo}


# ── Registrar pago ────────────────────────────────────────────────────────────────

@router.post("/{gasto_id}/pago")
def registrar_pago_gasto(
    gasto_id: int,
    data: PagoGastoCreate,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.id == gasto_id)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    g = q.first()
    if not g:
        raise HTTPException(404, "Gasto no encontrado")
    if g.estado_pago == "Pagado":
        raise HTTPException(400, "Este gasto ya está totalmente pagado")
    saldo_actual = float(g.saldo_pendiente) if g.saldo_pendiente is not None else float(g.monto or 0)
    if data.monto_pagado <= 0:
        raise HTTPException(400, "El monto a pagar debe ser mayor a 0")

    diferencia = round(data.monto_pagado - saldo_actual, 2)
    if diferencia > TOLERANCIA_REDONDEO + 0.01:
        raise HTTPException(400, f"El monto ({data.monto_pagado:.2f}) excede el saldo pendiente ({saldo_actual:.2f})")

    redondeo_tipo  = None
    redondeo_monto = 0.0
    if abs(diferencia) <= TOLERANCIA_REDONDEO + 0.01:
        nuevo_saldo  = 0.0
        nuevo_estado = "Pagado"
        if diferencia < -0.01:
            redondeo_tipo  = "ganancia"
            redondeo_monto = abs(diferencia)
        elif diferencia > 0.01:
            redondeo_tipo  = "perdida"
            redondeo_monto = diferencia
    else:
        nuevo_saldo  = round(saldo_actual - data.monto_pagado, 2)
        nuevo_estado = "Pago Parcial"

    pago = PagoGasto(
        gasto_id      = gasto_id,
        monto_pagado  = round(data.monto_pagado, 2),
        fecha_pago    = data.fecha_pago,
        metodo_pago   = data.metodo_pago,
        banco         = data.banco or None,
        numero_cuenta = data.numero_cuenta or None,
        numero_cheque = data.numero_cheque or None,
        numero_operacion = data.numero_operacion or None,
        created_at    = date.today(),
        redondeo_tipo  = redondeo_tipo,
        redondeo_monto = round(redondeo_monto, 2),
    )
    db.add(pago)

    g.saldo_pendiente = nuevo_saldo
    g.estado_pago      = nuevo_estado

    db.commit()
    db.refresh(g)
    return {
        "mensaje":         "Pago registrado",
        "saldo_pendiente": g.saldo_pendiente,
        "estado_pago":     g.estado_pago,
        "redondeo_tipo":   redondeo_tipo,
        "redondeo_monto":  round(redondeo_monto, 2),
    }


# ── Historial de pagos ────────────────────────────────────────────────────────────

@router.get("/{gasto_id}/historial-pagos")
def historial_pagos_gasto(
    gasto_id: int,
    db: Session = Depends(get_db),
    empresa_id: Optional[int] = Depends(get_empresa_id),
):
    q = db.query(Gasto).filter(Gasto.id == gasto_id)
    if empresa_id is not None:
        q = q.filter(Gasto.empresa_id == empresa_id)
    g = q.first()
    if not g:
        raise HTTPException(404, "Gasto no encontrado")
    pagos = db.query(PagoGasto).filter(PagoGasto.gasto_id == gasto_id).order_by(PagoGasto.fecha_pago.asc()).all()
    return [
        {
            "id":            p.id,
            "monto_pagado":  p.monto_pagado,
            "fecha_pago":    str(p.fecha_pago),
            "metodo_pago":   p.metodo_pago,
            "banco":         p.banco,
            "numero_cuenta": p.numero_cuenta,
            "numero_cheque": p.numero_cheque,
            "numero_operacion": p.numero_operacion,
            "created_at":    str(p.created_at) if p.created_at else None,
            "redondeo_tipo":  p.redondeo_tipo,
            "redondeo_monto": round(float(p.redondeo_monto or 0), 2),
        }
        for p in pagos
    ]
